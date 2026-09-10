
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import base64
import re
import calendar

import numpy as np
import pandas as pd
from pandas.io.formats.style import Styler
import openpyxl
import plotly.graph_objects as go
import requests
import streamlit as st


# =========================================================
# CONFIGURACIÓN GENERAL
# =========================================================
st.set_page_config(
    page_title="Centro de Control Gerencial VSE",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

ROYAL_BLUE = "#2563EB"
DARK_BLUE = "#60A5FA"
SECONDARY_BLUE = "#60A5FA"
LIGHT_BLUE = "#16263A"
BG = "#07111F"
CARD = "#101C2C"
PANEL = "#16263A"
TEXT = "#F8FAFC"
TEXT_2 = "#CBD5E1"
TEXT_3 = "#94A3B8"
GREEN = "#22C55E"
RED = "#EF4444"
AMBER = "#F59E0B"
ORANGE = "#F97316"
GRID = "#29415E"

st.markdown(
    f"""
    <style>
        :root {{
            --vse-bg: {BG};
            --vse-card: {CARD};
            --vse-panel: {PANEL};
            --vse-border: {GRID};
            --vse-text: {TEXT};
            --vse-muted: {TEXT_2};
            --vse-blue: {ROYAL_BLUE};
        }}

        .stApp {{
            background: {BG} !important;
            color: {TEXT} !important;
        }}
        html, body, [class*="css"] {{
            color: {TEXT};
        }}
        .block-container {{
            padding-top: 0.65rem;
            padding-bottom: 2.2rem;
            max-width: 1800px;
        }}
        h1, h2, h3, h4, h5, h6 {{
            color: {TEXT} !important;
            letter-spacing: -0.01em;
        }}
        p, label, small {{
            color: {TEXT_2};
        }}

        [data-testid="stSidebar"] {{
            background: #0B1626 !important;
            border-right: 1px solid {GRID} !important;
        }}
        [data-testid="stSidebar"] * {{
            color: #E8EEF7 !important;
        }}

        .hero {{
            background: linear-gradient(105deg, #0A2D59 0%, #114A95 48%, #2563EB 100%);
            padding: 19px 24px;
            border-radius: 14px;
            color: white;
            margin-bottom: 12px;
            box-shadow: 0 8px 26px rgba(0,0,0,.24);
            border: 1px solid rgba(96,165,250,.18);
        }}
        .hero-title {{
            font-size: 28px;
            font-weight: 850;
            margin: 0;
            color: #FFFFFF;
        }}
        .hero-sub {{
            font-size: 14px;
            font-weight: 600;
            margin-top: 5px;
            color: #DCE8F8;
        }}

        .base-status {{
            background: #0D1A2B;
            border: 1px solid {GRID};
            border-left: 4px solid {GREEN};
            border-radius: 10px;
            padding: 9px 13px;
            margin: 0 0 12px 0;
            color: {TEXT_2};
            font-size: 12px;
        }}
        .base-status strong {{ color: {TEXT}; }}

        .section-title {{
            background: #0F2D52;
            color: #FFFFFF;
            padding: 8px 12px;
            border-radius: 9px 9px 0 0;
            border: 1px solid {GRID};
            font-weight: 800;
            font-size: 13px;
            margin-top: 7px;
            letter-spacing: .25px;
        }}
        .subsection-tag {{
            color: {TEXT_2};
            font-size: 11px;
            font-weight: 800;
            letter-spacing: .7px;
            text-transform: uppercase;
            margin: 5px 0 7px 0;
        }}

        .kpi {{
            background: {CARD};
            border: 1px solid {GRID};
            border-left: 5px solid {ROYAL_BLUE};
            border-radius: 12px;
            padding: 13px 14px;
            min-height: 116px;
            box-shadow: 0 4px 14px rgba(0,0,0,.18);
        }}
        .kpi-label {{
            color: {TEXT_2};
            font-size: 11px;
            font-weight: 800;
            text-transform: uppercase;
            letter-spacing: .35px;
        }}
        .kpi-value {{
            color: {TEXT};
            font-size: 23px;
            font-weight: 850;
            margin-top: 6px;
        }}
        .kpi-delta {{
            margin-top: 7px;
            font-size: 11px;
            font-weight: 700;
        }}
        .kpi-note {{
            color: {TEXT_3};
            font-size: 10.5px;
            margin-top: 4px;
            line-height: 1.25;
        }}
        .small-note {{
            font-size: 11px;
            color: {TEXT_3};
        }}

        .fin-kpi {{
            background: {CARD};
            border: 1px solid {GRID};
            border-left: 5px solid {ROYAL_BLUE};
            border-radius: 12px;
            padding: 13px 14px;
            min-height: 116px;
            box-shadow: 0 4px 14px rgba(0,0,0,.18);
        }}
        .fin-kpi .label {{
            color: {TEXT_2};
            font-size: 11px;
            font-weight: 800;
            text-transform: uppercase;
            letter-spacing: .35px;
        }}
        .fin-kpi .value {{
            color: {TEXT};
            font-size: 23px;
            font-weight: 850;
            margin-top: 6px;
        }}
        .fin-kpi .note {{
            color: {TEXT_3};
            font-size: 10.5px;
            margin-top: 5px;
            line-height: 1.25;
        }}
        .fin-kpi.good {{ border-left-color: {GREEN}; }}
        .fin-kpi.warn {{ border-left-color: {AMBER}; }}
        .fin-kpi.bad {{ border-left-color: {RED}; }}

        .fin-panel {{
            background: {CARD};
            border: 1px solid {GRID};
            border-radius: 12px;
            padding: 13px 15px;
            margin: 7px 0 12px 0;
        }}
        .fin-panel strong {{ color: {TEXT}; }}
        .fin-note {{ color: {TEXT_2}; font-size: 11px; }}

        div[data-testid="stDataFrame"] {{
            background: {CARD} !important;
            border: 1px solid {GRID} !important;
            border-radius: 10px;
            overflow: hidden;
        }}

        div[data-baseweb="select"] > div,
        div[data-testid="stDateInput"] input,
        div[data-testid="stTextInput"] input,
        div[data-testid="stNumberInput"] input {{
            background: {CARD} !important;
            color: {TEXT} !important;
            border-color: {GRID} !important;
        }}
        div[data-baseweb="popover"] > div,
        ul[role="listbox"] {{
            background: {CARD} !important;
            color: {TEXT} !important;
        }}

        div[data-testid="stFileUploader"] section {{
            background: {CARD} !important;
            border-color: {GRID} !important;
        }}

        div[data-testid="stDownloadButton"] button,
        div[data-testid="stButton"] button {{
            background: #12335A;
            color: #FFFFFF !important;
            border: 1px solid #315A86;
            border-radius: 9px;
        }}
        div[data-testid="stDownloadButton"] button:hover,
        div[data-testid="stButton"] button:hover {{
            border-color: {SECONDARY_BLUE};
            background: #174776;
            color: #FFFFFF !important;
        }}

        div[data-testid="stSelectbox"] label,
        div[data-testid="stMultiSelect"] label,
        div[data-testid="stDateInput"] label,
        div[data-testid="stRadio"] label,
        div[data-testid="stFileUploader"] label {{
            font-weight: 700 !important;
            color: {TEXT_2} !important;
        }}

        [data-testid="stExpander"] {{
            background: #0D1A2B;
            border: 1px solid {GRID};
            border-radius: 10px;
        }}

        hr {{ border-color: {GRID}; }}
    </style>
    """,
    unsafe_allow_html=True,
)



# =========================================================
# SISTEMA VISUAL GLOBAL DE TABLAS · DARK NAVY EJECUTIVO
# =========================================================
st.markdown(
    """
    <style>
        .exec-table-wrap {
            width: 100%;
            overflow: auto;
            background: #101C2C;
            border: 1px solid #29415E;
            border-radius: 11px;
            margin: 4px 0 14px 0;
            box-shadow: 0 4px 14px rgba(0,0,0,.15);
        }

        .exec-table-wrap table {
            border-collapse: separate !important;
            border-spacing: 0 !important;
            width: 100% !important;
            min-width: 720px;
            margin: 0 !important;
            font-size: 11.5px !important;
            color: #F8FAFC !important;
        }

        .exec-table-wrap thead th {
            position: sticky;
            top: 0;
            z-index: 5;
            background: #12365E !important;
            color: #FFFFFF !important;
            font-weight: 800 !important;
            text-align: center !important;
            padding: 9px 10px !important;
            border-right: 1px solid #29415E !important;
            border-bottom: 1px solid #3B5C82 !important;
            white-space: nowrap !important;
        }

        .exec-table-wrap thead th:first-child {
            text-align: left !important;
        }

        .exec-table-wrap tbody td {
            background: #101C2C !important;
            color: #F8FAFC !important;
            padding: 8px 10px !important;
            border-right: 1px solid #20364F !important;
            border-bottom: 1px solid #20364F !important;
            text-align: right !important;
            white-space: nowrap !important;
            vertical-align: middle !important;
        }

        .exec-table-wrap tbody tr:nth-child(even) td {
            background: #13243A !important;
        }

        .exec-table-wrap tbody tr:hover td {
            background: #19334F !important;
        }

        .exec-table-wrap tbody td:first-child {
            text-align: left !important;
            font-weight: 650 !important;
            white-space: normal !important;
            min-width: 190px;
        }

        /* Acentos semánticos: discretos, no tipo semáforo */
        .exec-table-wrap td.cell-accent {
            color: #93C5FD !important;
            font-weight: 700 !important;
        }
        .exec-table-wrap td.cell-good {
            color: #86EFAC !important;
            font-weight: 700 !important;
        }
        .exec-table-wrap td.cell-warn {
            color: #FBBF24 !important;
            font-weight: 700 !important;
        }
        .exec-table-wrap td.cell-bad {
            color: #FCA5A5 !important;
            font-weight: 700 !important;
        }

        .exec-table-wrap::-webkit-scrollbar {
            height: 9px;
            width: 9px;
        }
        .exec-table-wrap::-webkit-scrollbar-track {
            background: #0B1626;
        }
        .exec-table-wrap::-webkit-scrollbar-thumb {
            background: #315A86;
            border-radius: 8px;
        }

        .exec-table-note {
            color: #94A3B8;
            font-size: 10px;
            margin-top: -7px;
            margin-bottom: 8px;
        }
    </style>
    """,
    unsafe_allow_html=True,
)


def _table_cell_classes(df: pd.DataFrame) -> pd.DataFrame:
    """
    Aplica acentos cromáticos discretos según el significado de la columna.
    El fondo permanece Dark Navy; sólo cambia el color tipográfico de celdas clave.
    """
    classes = pd.DataFrame("", index=df.index, columns=df.columns)

    for col in df.columns:
        label = str(col).strip().upper()

        # Pendientes / exposición: ámbar.
        if "PENDIENTE" in label:
            classes[col] = "cell-warn"

        # Porcentajes, participación, rentabilidad y variaciones: azul claro.
        elif (
            "%" in label
            or "RENTABILIDAD" in label
            or "PARTICIPACIÓN" in label
            or "PARTICIPACION" in label
            or "VARIACIÓN" in label
            or "VARIACION" in label
            or "COBERTURA" in label
        ):
            classes[col] = "cell-accent"

        # Dinero efectivamente realizado / margen: verde suave.
        elif (
            "RECAUDADO" in label
            or "PAGADO" in label
            or label == "MARGEN"
            or "CAJA OPERATIVA" in label
        ):
            classes[col] = "cell-good"

    # Valores negativos en columnas económicas relevantes: rojo suave.
    for col in df.columns:
        label = str(col).strip().upper()
        if any(k in label for k in ["MARGEN", "CAJA OPERATIVA", "VARIACIÓN", "VARIACION"]):
            numeric = pd.to_numeric(df[col], errors="coerce")
            neg_mask = numeric.lt(0)
            if neg_mask.any():
                classes.loc[neg_mask, col] = "cell-bad"

    # Estados operativos: acento semántico moderado.
    for col in df.columns:
        if str(col).strip().upper() in {"ESTADO", "ESTADO OP", "ESTADO OPERATIVO"}:
            values = df[col].astype("string").str.upper().fillna("")
            classes.loc[values.str.contains("CUMPLIDO", na=False), col] = "cell-good"
            classes.loc[values.str.contains("PROGRAM", na=False), col] = "cell-warn"
            classes.loc[values.str.contains("TRANSITO", na=False), col] = "cell-warn"

    return classes


def executive_dataframe(
    data,
    use_container_width=True,
    hide_index=False,
    height=None,
    **kwargs,
):
    """
    Sustituto visual de st.dataframe para todas las tablas del dashboard.
    Conserva los formatos de pandas Styler y añade:
    encabezado azul, filas dark navy, zebra suave, hover, scroll y contraste alto.
    """
    if isinstance(data, Styler):
        styler = data
        df_table = styler.data
    elif isinstance(data, pd.DataFrame):
        df_table = data
        styler = data.style
    else:
        # Compatibilidad defensiva.
        try:
            df_table = pd.DataFrame(data)
            styler = df_table.style
        except Exception:
            return st.markdown(str(data))

    # Ocultar índice para mantener estética ejecutiva.
    if hide_index:
        try:
            styler = styler.hide(axis="index")
        except Exception:
            pass

    # Clases semánticas generales.
    try:
        styler = styler.set_td_classes(_table_cell_classes(df_table))
    except Exception:
        pass

    # El header/filas se fuerzan por CSS global con !important.
    try:
        styler = styler.set_table_attributes('class="exec-table"')
    except Exception:
        pass

    html_table = styler.to_html()

    max_h = int(height) if isinstance(height, (int, float)) else 460
    max_h = max(160, min(max_h, 620))

    st.markdown(
        f'<div class="exec-table-wrap" style="max-height:{max_h}px;">{html_table}</div>',
        unsafe_allow_html=True,
    )


# Intercepta TODAS las tablas existentes sin reescribir 20 bloques por separado.
st.dataframe = executive_dataframe



# =========================================================
# SISTEMA VISUAL GLOBAL DE GRÁFICAS · ALTO CONTRASTE
# =========================================================
_original_plotly_chart = st.plotly_chart


def executive_plotly_chart(figure_or_data, *args, **kwargs):
    """
    Homologa TODAS las gráficas Plotly al Dark Navy ejecutivo.
    Fuerza contraste alto en leyendas, títulos, ejes, etiquetas y hover.
    """
    fig = figure_or_data

    try:
        fig.update_layout(
            paper_bgcolor="#07111F",
            plot_bgcolor="#101C2C",
            font=dict(
                color="#F8FAFC",
                family="Arial, sans-serif",
                size=12,
            ),
            title_font=dict(
                color="#F8FAFC",
                size=15,
            ),
            legend=dict(
                font=dict(
                    color="#F8FAFC",
                    size=12,
                ),
                bgcolor="rgba(7,17,31,0.92)",
                bordercolor="#29415E",
                borderwidth=1,
            ),
            hoverlabel=dict(
                bgcolor="#16263A",
                bordercolor="#3B5C82",
                font=dict(
                    color="#F8FAFC",
                    size=12,
                ),
            ),
        )

        fig.update_xaxes(
            tickfont=dict(color="#CBD5E1", size=11),
            title_font=dict(color="#CBD5E1", size=12),
            gridcolor="#29415E",
            zerolinecolor="#29415E",
            linecolor="#29415E",
        )
        fig.update_yaxes(
            tickfont=dict(color="#CBD5E1", size=11),
            title_font=dict(color="#CBD5E1", size=12),
            gridcolor="#29415E",
            zerolinecolor="#29415E",
            linecolor="#29415E",
        )

        # Refuerzo de ejes secundarios sin eliminar su configuración previa.
        for axis_name in ["yaxis2", "yaxis3", "xaxis2", "xaxis3"]:
            axis_obj = getattr(fig.layout, axis_name, None)
            if axis_obj is not None:
                axis_obj.update(
                    tickfont=dict(color="#CBD5E1", size=11),
                    title_font=dict(color="#CBD5E1", size=12),
                    gridcolor="#29415E",
                    zerolinecolor="#29415E",
                    linecolor="#29415E",
                )

        # Texto visible dentro de barras/pies cuando la traza lo soporta.
        for tr in fig.data:
            try:
                if hasattr(tr, "textfont"):
                    tr.textfont = dict(color="#F8FAFC")
            except Exception:
                pass

            # Leyendas y labels de pie/donut.
            try:
                if tr.type == "pie":
                    tr.insidetextfont = dict(color="#FFFFFF", size=12)
                    tr.outsidetextfont = dict(color="#F8FAFC", size=11)
            except Exception:
                pass

    except Exception:
        # El contraste nunca debe romper el dashboard.
        pass

    return _original_plotly_chart(fig, *args, **kwargs)


# Aplica automáticamente a todas las gráficas existentes.
st.plotly_chart = executive_plotly_chart



# =========================================================
# STORYTELLING GERENCIAL · JERARQUÍA Y LECTURA EN Z
# =========================================================
st.markdown(
    """
    <style>
        .story-banner {
            background: linear-gradient(110deg, #0D1A2B 0%, #112B49 100%);
            border: 1px solid #29415E;
            border-left: 5px solid #60A5FA;
            border-radius: 12px;
            padding: 13px 16px;
            margin: 6px 0 16px 0;
            box-shadow: 0 4px 14px rgba(0,0,0,.16);
        }
        .story-banner .eyebrow {
            color: #93C5FD;
            font-size: 10px;
            font-weight: 850;
            letter-spacing: .8px;
            text-transform: uppercase;
            margin-bottom: 5px;
        }
        .story-banner .story-text {
            color: #F8FAFC;
            font-size: 14px;
            font-weight: 650;
            line-height: 1.48;
        }
        .story-path {
            display: flex;
            align-items: center;
            justify-content: center;
            gap: 8px;
            flex-wrap: wrap;
            color: #CBD5E1;
            font-size: 10px;
            font-weight: 800;
            letter-spacing: .35px;
            text-transform: uppercase;
            margin: 2px 0 18px 0;
        }
        .story-path .node {
            background: #101C2C;
            border: 1px solid #29415E;
            border-radius: 18px;
            padding: 6px 10px;
            color: #DCE8F8;
        }
        .story-path .arrow {
            color: #60A5FA;
            font-size: 14px;
            font-weight: 900;
        }
        .question-tag {
            color: #93C5FD;
            font-size: 10px;
            font-weight: 850;
            text-transform: uppercase;
            letter-spacing: .8px;
            margin: 4px 0 5px 0;
        }
        .focus-grid {
            display: grid;
            grid-template-columns: repeat(5, minmax(150px, 1fr));
            gap: 10px;
            margin: 7px 0 16px 0;
        }
        .focus-card {
            background: #101C2C;
            border: 1px solid #29415E;
            border-top: 3px solid #F59E0B;
            border-radius: 11px;
            padding: 11px 12px;
            min-height: 100px;
        }
        .focus-card .label {
            color: #94A3B8;
            font-size: 9.5px;
            font-weight: 800;
            text-transform: uppercase;
            letter-spacing: .45px;
        }
        .focus-card .value {
            color: #F8FAFC;
            font-size: 15px;
            font-weight: 850;
            margin-top: 6px;
            line-height: 1.18;
        }
        .focus-card .note {
            color: #CBD5E1;
            font-size: 10px;
            margin-top: 5px;
            line-height: 1.25;
        }
        .story-divider {
            height: 1px;
            background: linear-gradient(90deg, transparent, #29415E, transparent);
            margin: 15px 0 17px 0;
        }
        @media (max-width: 1100px) {
            .focus-grid { grid-template-columns: repeat(2, 1fr); }
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# UTILIDADES
# =========================================================
MONTHS_ES = {
    1: "Ene", 2: "Feb", 3: "Mar", 4: "Abr", 5: "May", 6: "Jun",
    7: "Jul", 8: "Ago", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dic"
}

REQUIRED = [
    "CARGA", "FECHA CUM", "CLIENTE", "Quien Creó", "ESTADO OP",
    "TRAY. PROP", "LINEA NEG", "T.NEGOCIO",
    "V.CLIENTE", "V.CONDUCT", "T. VEHICULO",
    "PLACA", "CONDUCTO.1", "ESTADO FA", "FAC PROVE"
]


def norm_state(s: pd.Series) -> pd.Series:
    return (
        s.astype("string")
        .str.strip()
        .str.upper()
        .str.replace("Á", "A", regex=False)
        .str.replace("É", "E", regex=False)
        .str.replace("Í", "I", regex=False)
        .str.replace("Ó", "O", regex=False)
        .str.replace("Ú", "U", regex=False)
    )


@st.cache_data(show_spinner=False)
def load_excel(raw: bytes) -> pd.DataFrame:
    df = pd.read_excel(BytesIO(raw), sheet_name="TODOS", engine="openpyxl")
    df.columns = [str(c).strip() for c in df.columns]

    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        raise ValueError("Faltan columnas requeridas: " + ", ".join(missing))

    # Fecha maestra: CARGA.
    df["CARGA"] = pd.to_datetime(df["CARGA"], errors="coerce")
    df["FECHA_CARGA"] = df["CARGA"].dt.normalize()

    for c in ["FECHA R", "F PRECUMP", "FECHA CUM"]:
        if c in df.columns:
            df[c] = pd.to_datetime(df[c], errors="coerce")

    for c in ["V.CLIENTE", "V.CONDUCT", "TRAY. PROP"]:
        df[c] = pd.to_numeric(df[c], errors="coerce")

    df["ESTADO OP N"] = norm_state(df["ESTADO OP"])
    df["Tipo Flota"] = np.select(
        [df["TRAY. PROP"].eq(1), df["TRAY. PROP"].eq(0)],
        ["FLOTA PROPIA", "TERCEROS"],
        default="SIN CLASIFICAR",
    )

    # Tipología: campo T. VEHICULO.
    df["TIPOLOGIA"] = (
        df["T. VEHICULO"]
        .astype("string")
        .fillna("SIN TIPOLOGÍA")
        .str.strip()
    )

    # Extrae capacidad numérica cuando el nombre contiene "XX PASAJEROS".
    cap = (
        df["TIPOLOGIA"]
        .str.extract(r"(\d+)\s*PASAJER", flags=re.IGNORECASE, expand=False)
    )
    df["CAPACIDAD_PAX"] = pd.to_numeric(cap, errors="coerce")

    # Conductor oficial: columna W en el archivo = CONDUCTO.1
    df["CONDUCTOR_NOMBRE"] = (
        df["CONDUCTO.1"].astype("string").fillna("SIN CONDUCTOR").str.strip()
    )

    df["PLACA"] = (
        df["PLACA"]
        .astype("string")
        .fillna("SIN PLACA")
        .str.strip()
        .str.upper()
    )
    df["CLIENTE"] = (
        df["CLIENTE"]
        .astype("string")
        .fillna("SIN CLIENTE")
        .str.strip()
        .str.upper()
    )
    df["Quien Creó"] = df["Quien Creó"].astype("string").fillna("SIN COORDINADOR").str.strip()

    df["AÑO"] = df["FECHA_CARGA"].dt.year
    df["MES_NUM"] = df["FECHA_CARGA"].dt.month
    df["MES"] = df["MES_NUM"].map(MONTHS_ES)
    df["BIMESTRE_NUM"] = ((df["MES_NUM"] - 1) // 2 + 1).astype("Int64")
    df["TRIMESTRE_NUM"] = ((df["MES_NUM"] - 1) // 3 + 1).astype("Int64")
    df["SEMESTRE_NUM"] = np.where(df["MES_NUM"] <= 6, 1, 2)

    return df


def valid_services(df: pd.DataFrame) -> pd.DataFrame:
    return df.loc[df["ESTADO OP N"].ne("ANULADO")].copy()


def fmt_int(x) -> str:
    if pd.isna(x):
        return "0"
    return f"{int(round(x)):,}".replace(",", ".")


def fmt_money(x) -> str:
    """COP completo y fácil de leer: $7.020.000.000."""
    if pd.isna(x):
        return "$0"
    return "$" + f"{int(round(float(x))):,}".replace(",", ".")


def fmt_pct(x) -> str:
    if pd.isna(x):
        return "—"
    return f"{x*100:.2f}%".replace(".", ",")


def delta_html(delta, good_when_up=True, suffix="%"):
    if delta is None or pd.isna(delta):
        return ""
    favorable = delta >= 0 if good_when_up else delta <= 0
    color = GREEN if favorable else RED
    arrow = "▲" if delta > 0 else ("▼" if delta < 0 else "▬")

    # Tanto porcentajes como puntos porcentuales parten de una razón decimal.
    if suffix in ("%", " p.p."):
        val = abs(delta) * 100
    else:
        val = abs(delta)

    label = f"{val:.2f}{suffix}".replace(".", ",")
    return (
        f"<span style='color:{color}'>{arrow} {label}</span> "
        f"<span style='color:{TEXT_2}'>último periodo seleccionado vs anterior</span>"
    )


def kpi_card(label, value, delta=None, good_when_up=True, note="", suffix="%", show_delta=True):
    d = delta_html(delta, good_when_up, suffix) if show_delta else ""
    delta_block = f'<div class="kpi-delta">{d}</div>' if (show_delta and d) else ""
    note_block = f'<div class="kpi-note">{note}</div>' if note else ""

    html = (
        '<div class="kpi">'
        f'<div class="kpi-label">{label}</div>'
        f'<div class="kpi-value">{value}</div>'
        f'{delta_block}'
        f'{note_block}'
        '</div>'
    )
    st.markdown(html, unsafe_allow_html=True)


def safe_div(a, b):
    return a / b if b not in (0, None) and not pd.isna(b) else np.nan


def aggregate_period(df: pd.DataFrame, grain: str) -> pd.DataFrame:
    x = df.copy()
    if grain == "Mensual":
        x["PERIODO_ORDEN"] = x["AÑO"] * 100 + x["MES_NUM"]
        x["PERIODO"] = x.apply(lambda r: f"{MONTHS_ES.get(r['MES_NUM'], '')} {int(r['AÑO'])}", axis=1)
    elif grain == "Bimestral":
        x["PERIODO_ORDEN"] = x["AÑO"] * 10 + x["BIMESTRE_NUM"].astype(int)
        x["PERIODO"] = x.apply(lambda r: f"B{int(r['BIMESTRE_NUM'])} {int(r['AÑO'])}", axis=1)
    elif grain == "Trimestral":
        x["PERIODO_ORDEN"] = x["AÑO"] * 10 + x["TRIMESTRE_NUM"].astype(int)
        x["PERIODO"] = x.apply(lambda r: f"T{int(r['TRIMESTRE_NUM'])} {int(r['AÑO'])}", axis=1)
    else:
        x["PERIODO_ORDEN"] = x["AÑO"] * 10 + x["SEMESTRE_NUM"].astype(int)
        x["PERIODO"] = x.apply(lambda r: f"S{int(r['SEMESTRE_NUM'])} {int(r['AÑO'])}", axis=1)

    g = (
        x.groupby(["PERIODO_ORDEN", "PERIODO"], dropna=False)
        .agg(
            Servicios=("CLIENTE", "size"),
            Facturacion=("V.CLIENTE", "sum"),
            Costos=("V.CONDUCT", "sum"),
        )
        .reset_index()
        .sort_values("PERIODO_ORDEN")
    )
    g["Margen"] = g["Facturacion"] - g["Costos"]
    g["Rentabilidad"] = np.where(g["Facturacion"].ne(0), g["Margen"] / g["Facturacion"], np.nan)
    g["Var_Servicios"] = g["Servicios"].pct_change()
    g["Var_Facturacion"] = g["Facturacion"].pct_change()
    g["Var_Margen"] = g["Margen"].pct_change()
    g["Var_Rent_PP"] = g["Rentabilidad"].diff()

    # Protección estadística:
    # si el último periodo está incompleto, NO se calcula una variación contra
    # un periodo anterior completo, porque sería una comparación no homogénea.
    max_date = x["FECHA_CARGA"].max()
    periodo_incompleto = False

    if pd.notna(max_date):
        year = int(max_date.year)
        month = int(max_date.month)

        if grain == "Mensual":
            last_day = calendar.monthrange(year, month)[1]
            periodo_incompleto = int(max_date.day) < last_day

        elif grain == "Bimestral":
            end_month = ((month - 1) // 2 + 1) * 2
            last_day = calendar.monthrange(year, end_month)[1]
            periodo_incompleto = not (
                month == end_month and int(max_date.day) == last_day
            )

        elif grain == "Trimestral":
            end_month = ((month - 1) // 3 + 1) * 3
            last_day = calendar.monthrange(year, end_month)[1]
            periodo_incompleto = not (
                month == end_month and int(max_date.day) == last_day
            )

        else:
            end_month = 6 if month <= 6 else 12
            last_day = calendar.monthrange(year, end_month)[1]
            periodo_incompleto = not (
                month == end_month and int(max_date.day) == last_day
            )

    g["Periodo_Incompleto"] = False
    if len(g) and periodo_incompleto:
        idx = g.index[-1]
        g.loc[idx, "Periodo_Incompleto"] = True
        g.loc[idx, ["Var_Servicios", "Var_Facturacion", "Var_Margen", "Var_Rent_PP"]] = np.nan

    return g


def trend_chart(g: pd.DataFrame, value_col: str, var_col: str, title: str, money=False):
    vals = g[value_col]
    labels = [fmt_money(v) if money else fmt_int(v) for v in vals]

    var_real_pct = g[var_col].astype(float) * 100
    visual_cap = 200.0
    var_visual_pct = var_real_pct.clip(-visual_cap, visual_cap)

    text = []
    custom = []
    for v in var_real_pct:
        if pd.isna(v):
            text.append("")
            custom.append(["Sin comparación homogénea"])
        else:
            arrow = "▲" if v > 0 else ("▼" if v < 0 else "▬")
            marker = " *" if abs(v) > visual_cap else ""
            text.append(f"{arrow} {abs(v):.2f}%{marker}")
            custom.append([f"{v:.2f}%"])

    fig = go.Figure()
    fig.add_trace(
        go.Bar(
            x=g["PERIODO"],
            y=vals,
            name=title,
            marker_color=ROYAL_BLUE,
            text=labels,
            textposition="outside",
            cliponaxis=False,
            hovertemplate="%{x}<br>%{text}<extra></extra>",
        )
    )
    fig.add_trace(
        go.Scatter(
            x=g["PERIODO"],
            y=var_visual_pct,
            name="Variación %",
            mode="lines+markers+text",
            line=dict(color=SECONDARY_BLUE, width=2),
            marker=dict(size=8),
            text=text,
            customdata=np.array(custom, dtype=object),
            textposition="top center",
            yaxis="y2",
            hovertemplate="%{x}<br>Variación real: %{customdata[0]}<extra></extra>",
        )
    )

    finite = var_visual_pct.dropna()
    if len(finite):
        limit = max(30.0, min(visual_cap, float(np.nanmax(np.abs(finite))) * 1.2))
    else:
        limit = 30.0

    fig.update_layout(
        height=360,
        margin=dict(l=20, r=30, t=40, b=20),
        plot_bgcolor="#101C2C",
        paper_bgcolor="#07111F",
        font=dict(color=TEXT),
        legend=dict(orientation="h", y=1.16, x=0, font=dict(color="#F8FAFC", size=12)),
        yaxis=dict(gridcolor=GRID, zeroline=False, title="COP" if money else "Servicios"),
        yaxis2=dict(
            overlaying="y",
            side="right",
            title="Variación %",
            ticksuffix="%",
            showgrid=False,
            range=[-limit, limit],
        ),
        hovermode="x unified",
    )
    return fig


def top_table(df, group, value="Servicios", n=7):
    v = valid_services(df)
    g = (
        v.groupby(group, dropna=False)
        .agg(
            Servicios=("CLIENTE", "size"),
            Facturacion=("V.CLIENTE", "sum"),
            Costos=("V.CONDUCT", "sum"),
        )
        .reset_index()
    )
    g["Margen"] = g["Facturacion"] - g["Costos"]
    g["Rentabilidad %"] = np.where(g["Facturacion"].ne(0), g["Margen"] / g["Facturacion"], np.nan)
    if value == "Facturacion":
        g = g.sort_values("Facturacion", ascending=False)
    elif value == "Margen":
        g = g.sort_values("Margen", ascending=False)
    else:
        g = g.sort_values("Servicios", ascending=False)
    return g.head(n)


# =========================================================
# NAVEGACIÓN PRINCIPAL
# =========================================================
with st.sidebar:
    st.markdown("### 🧭 Navegación")
    page_mode = st.radio(
        "Módulo",
        ["CENTRO DE CONTROL", "POSICIÓN FINANCIERA Y PROYECCIONES"],
        index=0,
        label_visibility="collapsed",
        key="main_navigation",
    )

if page_mode == "CENTRO DE CONTROL":
    st.markdown(
        """
        <div class="hero">
            <div class="hero-title">📊 CENTRO DE CONTROL GERENCIAL VSE</div>
            <div class="hero-sub">Qué pasó · Cómo cerramos · Dónde está el problema · Quién lo explica · Dónde actuar</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
else:
    st.markdown(
        """
        <div class="hero">
            <div class="hero-title">💹 POSICIÓN FINANCIERA Y PROYECCIONES</div>
            <div class="hero-sub">Generamos · Cobramos · Pagamos · Medimos riesgo · Proyectamos</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

# =========================================================
# BASE VIGENTE · CARGA AUTOMÁTICA + ACTUALIZACIÓN COMPARTIDA
# =========================================================
MASTER_FILE = Path("TRAYECTOS TODOS__2026.xlsx")
BOGOTA_TZ = ZoneInfo("America/Bogota")


def _secret(name: str, default: str = "") -> str:
    """Lee Streamlit Secrets sin impedir el arranque si aún no están configurados."""
    try:
        value = st.secrets.get(name, default)
        return str(value).strip() if value is not None else default
    except Exception:
        return default


GITHUB_TOKEN = _secret("GITHUB_TOKEN")
GITHUB_REPO = _secret("GITHUB_REPO")  # Ej.: empresa/SEGUIMIENTO_OP_VSE
GITHUB_BRANCH = _secret("GITHUB_BRANCH", "main")
GITHUB_MASTER_PATH = _secret("GITHUB_MASTER_PATH", MASTER_FILE.name)

GITHUB_SHARED_READY = bool(GITHUB_TOKEN and GITHUB_REPO and GITHUB_MASTER_PATH)


def github_headers() -> dict:
    return {
        "Authorization": f"Bearer {GITHUB_TOKEN}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }


def fetch_master_from_github():
    """Descarga el Excel maestro como binario RAW autenticado."""
    if not GITHUB_SHARED_READY:
        return None

    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{GITHUB_MASTER_PATH}"

    # Metadata para verificar que existe y obtener SHA.
    meta_headers = github_headers()
    meta_headers["Accept"] = "application/vnd.github+json"
    meta_response = requests.get(
        url, headers=meta_headers, params={"ref": GITHUB_BRANCH}, timeout=25
    )
    if meta_response.status_code != 200:
        return None

    meta = meta_response.json()

    # Para Excel de varios MB no usamos el campo content/base64 del JSON.
    # Pedimos directamente los bytes RAW reales del archivo.
    raw_headers = github_headers()
    raw_headers["Accept"] = "application/vnd.github.raw+json"
    raw_response = requests.get(
        url, headers=raw_headers, params={"ref": GITHUB_BRANCH}, timeout=60
    )
    if raw_response.status_code != 200:
        return None

    raw_file = raw_response.content
    if not raw_file:
        return None

    if raw_file.startswith(b"version https://git-lfs.github.com/spec"):
        raise ValueError(
            "El archivo de GitHub es un puntero Git LFS, no el Excel real."
        )

    # Los .xlsx son contenedores ZIP y comienzan normalmente por PK.
    if not raw_file.startswith(b"PK"):
        raise ValueError(
            "GitHub no devolvió un archivo .xlsx válido."
        )

    updated_at = None
    commits_url = f"https://api.github.com/repos/{GITHUB_REPO}/commits"
    commits = requests.get(
        commits_url,
        headers=github_headers(),
        params={"path": GITHUB_MASTER_PATH, "sha": GITHUB_BRANCH, "per_page": 1},
        timeout=20,
    )
    if commits.status_code == 200 and commits.json():
        iso_date = commits.json()[0].get("commit", {}).get("committer", {}).get("date")
        if iso_date:
            updated_at = (
                datetime.fromisoformat(iso_date.replace("Z", "+00:00"))
                .astimezone(BOGOTA_TZ)
            )

    return {
        "raw": raw_file,
        "sha": meta.get("sha"),
        "name": Path(GITHUB_MASTER_PATH).name,
        "source": "GitHub compartido",
        "updated_at": updated_at,
    }

def publish_master_to_github(raw_file: bytes, original_name: str):
    """Publica la nueva base como archivo maestro compartido."""
    if not GITHUB_SHARED_READY:
        return False, "La persistencia compartida de GitHub no está configurada."

    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{GITHUB_MASTER_PATH}"
    try:
        current = requests.get(
            url,
            headers=github_headers(),
            params={"ref": GITHUB_BRANCH},
            timeout=25,
        )

        payload = {
            "message": (
                "Actualiza base maestra VSE desde dashboard · "
                + datetime.now(BOGOTA_TZ).strftime("%Y-%m-%d %H:%M")
            ),
            "content": base64.b64encode(raw_file).decode("ascii"),
            "branch": GITHUB_BRANCH,
        }

        if current.status_code == 200:
            payload["sha"] = current.json().get("sha")
        elif current.status_code not in (404,):
            return False, f"GitHub respondió {current.status_code} al consultar la base vigente."

        update = requests.put(
            url,
            headers=github_headers(),
            json=payload,
            timeout=45,
        )

        if update.status_code not in (200, 201):
            detail = update.json().get("message", update.text[:200])
            return False, f"No fue posible publicar la base: {detail}"

        return True, "Base maestra publicada correctamente."
    except Exception as exc:
        return False, f"No fue posible conectar con GitHub: {exc}"


# 1. La app abre SIEMPRE con la última base vigente.
master_info = None

# En la sesión que acaba de publicar, usar inmediatamente la nueva versión.
if st.session_state.get("master_override_raw"):
    master_info = {
        "raw": st.session_state["master_override_raw"],
        "name": MASTER_FILE.name,
        "source": "Actualización recién publicada",
        "updated_at": st.session_state.get("master_override_updated_at"),
    }
else:
    # Si hay Secrets configurados, consultar GitHub en tiempo real.
    master_info = fetch_master_from_github()

# Respaldo 1: archivo maestro con nombre esperado en el repositorio/despliegue.
if master_info is None and MASTER_FILE.exists():
    master_info = {
        "raw": MASTER_FILE.read_bytes(),
        "name": MASTER_FILE.name,
        "source": "Base incluida en el despliegue",
        "updated_at": None,
    }

# Respaldo 2: autodetectar cualquier Excel local que contenga la hoja TODOS.
# Esto evita que el dashboard quede vacío si el archivo fue subido con otro nombre.
if master_info is None:
    local_candidates = sorted(
        [p for p in Path(".").glob("*.xlsx") if not p.name.startswith("~$")],
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    for candidate in local_candidates:
        try:
            candidate_raw = candidate.read_bytes()
            # Validar estructura con la misma función oficial del dashboard.
            _candidate_df = load_excel(candidate_raw)
            if _candidate_df["FECHA_CARGA"].notna().sum() > 0:
                master_info = {
                    "raw": candidate_raw,
                    "name": candidate.name,
                    "source": "Base Excel autodetectada en el repositorio",
                    "updated_at": datetime.fromtimestamp(
                        candidate.stat().st_mtime, tz=BOGOTA_TZ
                    ),
                }
                break
        except Exception:
            continue

# Respaldo 3: si NO existe ninguna base persistente, permitir carga temporal
# en lugar de dejar una pantalla de error sin salida.
if master_info is None:
    st.warning(
        "⚠️ No se encontró una base maestra persistente en el despliegue. "
        "Puedes cargar una base ahora para visualizar el dashboard en esta sesión."
    )
    emergency_upload = st.file_uploader(
        "Cargar base temporal para iniciar",
        type=["xlsx"],
        accept_multiple_files=False,
        key="emergency_master_upload",
        help="La hoja debe llamarse TODOS y conservar las columnas requeridas.",
    )

    if emergency_upload is None:
        st.info(
            "Para que el dashboard abra automáticamente para todos, deja al menos un archivo Excel "
            "válido en el repositorio o configura la persistencia compartida desde Streamlit Secrets."
        )
        st.stop()

    emergency_raw = emergency_upload.getvalue()
    try:
        emergency_df = load_excel(emergency_raw)
        if emergency_df["FECHA_CARGA"].notna().sum() == 0:
            raise ValueError("La columna CARGA no contiene fechas válidas.")
        master_info = {
            "raw": emergency_raw,
            "name": emergency_upload.name,
            "source": "Base temporal cargada en esta sesión",
            "updated_at": datetime.now(BOGOTA_TZ),
        }
    except Exception as exc:
        st.error(f"❌ La base temporal no es válida: {exc}")
        st.stop()

raw = master_info["raw"]

try:
    df_all = load_excel(raw)
except ValueError as e:
    st.error(f"❌ La base vigente tiene una estructura no válida: {e}")
    st.stop()
except Exception as e:
    st.error(f"❌ No fue posible leer la base vigente: {e}")
    st.stop()

if df_all["FECHA_CARGA"].notna().sum() == 0:
    st.error("❌ La columna CARGA de la base vigente no contiene fechas válidas.")
    st.stop()

data_cutoff = df_all["FECHA_CARGA"].max()
data_start = df_all["FECHA_CARGA"].min()
valid_rows = len(df_all)
valid_dates = int(df_all["FECHA_CARGA"].notna().sum())

updated_label = (
    master_info["updated_at"].strftime("%d/%m/%Y %H:%M")
    if master_info.get("updated_at")
    else "versión del despliegue"
)

# Trazabilidad visible sin bloquear el dashboard.
st.markdown(
    f"""
    <div class="base-status">
        <strong>Base vigente:</strong> {master_info["name"]}
        &nbsp;·&nbsp; <strong>Registros:</strong> {fmt_int(valid_rows)}
        &nbsp;·&nbsp; <strong>Corte:</strong> {data_cutoff.strftime('%d/%m/%Y') if pd.notna(data_cutoff) else 'N/D'}
        &nbsp;·&nbsp; <strong>Actualización:</strong> {updated_label}
        &nbsp;·&nbsp; <strong>Origen:</strong> {master_info["source"]}
    </div>
    """,
    unsafe_allow_html=True,
)

# 2. Actualización opcional: no bloquea la visualización.
with st.sidebar:
    st.markdown("### 📁 Base vigente")
    st.success("✅ Dashboard disponible")
    st.caption(f"Corte actual: {data_cutoff.strftime('%d/%m/%Y') if pd.notna(data_cutoff) else 'N/D'}")
    st.caption(f"Registros: {fmt_int(valid_rows)}")

    with st.expander("🔄 Actualizar base vigente", expanded=False):
        st.caption(
            "Carga una nueva base únicamente cuando quieras reemplazar la información visible para todos."
        )

        uploaded_file = st.file_uploader(
            "Nueva base Excel",
            type=["xlsx"],
            accept_multiple_files=False,
            help="El nombre puede variar. La hoja debe llamarse TODOS.",
            key="shared_master_uploader",
        )

        # Plantilla vacía de estructura.
        template_cols = REQUIRED.copy()
        for extra_col in ["FECHA R", "F PRECUMP"]:
            if extra_col not in template_cols:
                template_cols.append(extra_col)

        template_buffer = BytesIO()
        with pd.ExcelWriter(template_buffer, engine="openpyxl") as writer:
            pd.DataFrame(columns=template_cols).to_excel(
                writer, sheet_name="TODOS", index=False
            )
        template_buffer.seek(0)

        st.download_button(
            "⬇️ Plantilla de estructura",
            data=template_buffer.getvalue(),
            file_name="PLANTILLA_BASE_VSE.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

        if uploaded_file is not None:
            candidate_raw = uploaded_file.getvalue()

            try:
                candidate_df = load_excel(candidate_raw)
                if candidate_df["FECHA_CARGA"].notna().sum() == 0:
                    raise ValueError("La columna CARGA no contiene fechas válidas.")

                candidate_cutoff = candidate_df["FECHA_CARGA"].max()
                candidate_start = candidate_df["FECHA_CARGA"].min()

                st.success("✓ Estructura validada")
                st.caption(f"Archivo: {uploaded_file.name}")
                st.caption(f"Registros: {fmt_int(len(candidate_df))}")
                if pd.notna(candidate_start) and pd.notna(candidate_cutoff):
                    st.caption(
                        "Periodo: "
                        f"{candidate_start.strftime('%d/%m/%Y')} — "
                        f"{candidate_cutoff.strftime('%d/%m/%Y')}"
                    )

                if not GITHUB_SHARED_READY:
                    st.warning(
                        "Para que esta actualización quede visible para todos, "
                        "configura GITHUB_TOKEN, GITHUB_REPO, GITHUB_BRANCH y GITHUB_MASTER_PATH "
                        "en Streamlit Secrets."
                    )

                publish_clicked = st.button(
                    "✅ Publicar como base vigente",
                    use_container_width=True,
                    disabled=not GITHUB_SHARED_READY,
                    key="publish_shared_master",
                )

                if publish_clicked:
                    with st.spinner("Publicando y actualizando el dashboard..."):
                        ok, message = publish_master_to_github(
                            candidate_raw, uploaded_file.name
                        )
                    if ok:
                        load_excel.clear()
                        st.session_state["master_override_raw"] = candidate_raw
                        st.session_state["master_override_updated_at"] = datetime.now(BOGOTA_TZ)
                        st.success("✅ Nueva base publicada. El dashboard se actualizará ahora.")
                        st.rerun()
                    else:
                        st.error(message)

            except Exception as e:
                st.error(f"❌ Base no válida: {e}")
                st.caption(
                    "No se publicará ningún cambio hasta que la estructura sea correcta."
                )


# =========================================================
# FILTROS
# =========================================================
with st.sidebar:
    st.markdown("### 🔎 Filtros principales")

    grain = st.selectbox(
        "Periodo de análisis",
        ["Mensual", "Bimestral", "Trimestral", "Semestral"],
        index=0,
    )

    years = sorted(df_all["AÑO"].dropna().astype(int).unique().tolist())
    selected_years = st.multiselect("Año", years, default=years)

    min_date = df_all["FECHA_CARGA"].min().date()
    max_date = df_all["FECHA_CARGA"].max().date()
    date_range = st.date_input(
        "Rango de fechas (CARGA)",
        value=(min_date, max_date),
        min_value=min_date,
        max_value=max_date,
    )
    if isinstance(date_range, tuple) and len(date_range) == 2:
        start_date, end_date = date_range
    else:
        start_date, end_date = min_date, max_date

    # Meses en español, seleccionables de forma múltiple.
    # Siempre se calculan desde FECHA_CARGA.
    MESES_ES = {
        1: "ENERO",
        2: "FEBRERO",
        3: "MARZO",
        4: "ABRIL",
        5: "MAYO",
        6: "JUNIO",
        7: "JULIO",
        8: "AGOSTO",
        9: "SEPTIEMBRE",
        10: "OCTUBRE",
        11: "NOVIEMBRE",
        12: "DICIEMBRE",
    }
    meses_disponibles_num = sorted(
        df_all["FECHA_CARGA"].dropna().dt.month.unique().tolist()
    )
    meses_disponibles = [
        MESES_ES[int(m)] for m in meses_disponibles_num
    ]

    selected_months = st.multiselect(
        "Mes(es)",
        options=meses_disponibles,
        default=[],
        placeholder="Todos los meses",
        help="Selecciona uno o varios meses. El filtro usa la fecha CARGA.",
    )

    def multi_filter(label, col):
        opts = sorted([x for x in df_all[col].dropna().astype(str).unique().tolist() if x and x != "<NA>"])
        return st.multiselect(label, opts, default=[])

    f_client = multi_filter("Cliente", "CLIENTE")
    f_coord = multi_filter("Quién creó (Coordinador)", "Quien Creó")
    f_estado = multi_filter("Estado Operativo", "ESTADO OP N")
    f_flota = multi_filter("Tipo Flota", "Tipo Flota")
    f_linea = multi_filter("Línea de Negocio", "LINEA NEG")
    f_negocio = multi_filter("Tipo de Negocio", "T.NEGOCIO")
    f_tipologia = multi_filter("Tipología / Tipo Vehículo", "TIPOLOGIA")

    with st.expander("Filtros avanzados", expanded=False):
        f_placa = multi_filter("Placa", "PLACA")
        f_conductor = multi_filter("Conductor", "CONDUCTOR_NOMBRE")

    if st.button("↻ Limpiar filtros", use_container_width=True):
        st.rerun()


def apply_filters(df):
    x = df.copy()
    x = x[x["AÑO"].isin(selected_years)] if selected_years else x.iloc[0:0]
    x = x[
        x["FECHA_CARGA"].between(
            pd.Timestamp(start_date),
            pd.Timestamp(end_date) + pd.Timedelta(days=1) - pd.Timedelta(microseconds=1),
        )
    ]

    # Filtro múltiple de meses sobre CARGA.
    if selected_months:
        selected_month_numbers = [
            num for num, nombre in MESES_ES.items()
            if nombre in selected_months
        ]
        x = x[x["FECHA_CARGA"].dt.month.isin(selected_month_numbers)]

    mapping = [
        ("CLIENTE", f_client),
        ("Quien Creó", f_coord),
        ("ESTADO OP N", f_estado),
        ("Tipo Flota", f_flota),
        ("LINEA NEG", f_linea),
        ("T.NEGOCIO", f_negocio),
        ("TIPOLOGIA", f_tipologia),
        ("PLACA", f_placa),
        ("CONDUCTOR_NOMBRE", f_conductor),
    ]
    for col, values in mapping:
        if values:
            x = x[x[col].astype(str).isin(values)]
    return x


df = apply_filters(df_all)
valid = valid_services(df)

if valid.empty:
    st.warning("No hay registros válidos para los filtros seleccionados.")
    st.stop()


# =========================================================
# PÁGINA 2 · POSICIÓN FINANCIERA Y PROYECCIONES
# =========================================================
if page_mode == "POSICIÓN FINANCIERA Y PROYECCIONES":
    st.markdown(
        """
        <style>
            .stApp { background: #07111F !important; color: #FFFFFF !important; }
            .block-container { max-width: 1800px; }
            h1, h2, h3, h4 { color: #FFFFFF !important; }
            p, span, label { color: #E2E8F0; }
            [data-testid="stSidebar"] { background: #0B1626 !important; border-right: 1px solid #263B55 !important; }
            [data-testid="stSidebar"] * { color: #E8EEF7 !important; }
            .hero { background: linear-gradient(90deg, #0A2D59 0%, #064AA3 58%, #0B5ED7 100%) !important; box-shadow: 0 6px 22px rgba(0,0,0,.25) !important; }
            .section-title { background: #0F2D52 !important; color: #FFFFFF !important; border: 1px solid #263B55; }
            .fin-kpi { background: #0F1C2E; border: 1px solid #263B55; border-left: 5px solid #1769AA; border-radius: 12px; padding: 13px 14px; min-height: 118px; box-shadow: 0 3px 12px rgba(0,0,0,.20); }
            .fin-kpi .label { color: #CBD5E1; font-size: 11px; font-weight: 800; text-transform: uppercase; letter-spacing: .35px; }
            .fin-kpi .value { color: #FFFFFF; font-size: 23px; font-weight: 850; margin-top: 6px; }
            .fin-kpi .note { color: #94A3B8; font-size: 10.5px; margin-top: 5px; line-height: 1.25; }
            .fin-kpi.good { border-left-color: #2DA66F; }
            .fin-kpi.warn { border-left-color: #D98E04; }
            .fin-kpi.bad { border-left-color: #C94B4B; }
            .fin-panel { background: #0F1C2E; border: 1px solid #263B55; border-radius: 12px; padding: 14px 16px; margin: 6px 0 12px 0; }
            .fin-panel strong { color: #FFFFFF; }
            .fin-note { color:#CBD5E1; font-size: 11px; }
            div[data-testid="stDataFrame"] { border: 1px solid #263B55 !important; border-radius: 10px; overflow: hidden; }
            div[data-testid="stDownloadButton"] button { background:#0F2D52; color:white; border:1px solid #2E5A89; }
            div[data-testid="stDownloadButton"] button:hover { border-color:#64B5F6; color:white; }
            div[data-testid="stSelectbox"] label,
            div[data-testid="stMultiSelect"] label,
            div[data-testid="stDateInput"] label,
            div[data-testid="stRadio"] label { color:#E2E8F0 !important; }
        </style>
        """,
        unsafe_allow_html=True,
    )

    def fin_card(label, value, note="", tone=""):
        tone_class = f" {tone}" if tone else ""
        html = (
            f'<div class="fin-kpi{tone_class}">'
            f'<div class="label">{label}</div>'
            f'<div class="value">{value}</div>'
            f'<div class="note">{note}</div>'
            '</div>'
        )
        st.markdown(html, unsafe_allow_html=True)

    def norm_fin_state(series: pd.Series) -> pd.Series:
        return norm_state(series.fillna(""))

    fin = valid.copy()
    fin["ESTADO_FA_N"] = norm_fin_state(fin["ESTADO FA"])
    fin["FAC_PROVE_N"] = norm_fin_state(fin["FAC PROVE"])

    fin["CLAS_CLIENTE"] = np.where(
        fin["ESTADO_FA_N"].eq("FACT DEFINITIVA"),
        "RECAUDADO",
        "PENDIENTE POR RECAUDAR",
    )

    # Compatibilidad histórica: la base local usa PAGADO; la nueva nomenclatura puede usar FACTURADO.
    paid_supplier_states = {"PAGADO", "FACTURADO"}
    fin["CLAS_PROVE"] = np.select(
        [fin["FAC_PROVE_N"].isin(paid_supplier_states), fin["FAC_PROVE_N"].eq("EN ESPERA")],
        ["PAGADO A TERCEROS", "PENDIENTE POR PAGAR"],
        default="POR CLASIFICAR",
    )

    servicios_fin = len(fin)
    produccion_fin = float(fin["V.CLIENTE"].fillna(0).sum())
    costo_fin = float(fin["V.CONDUCT"].fillna(0).sum())
    margen_fin = produccion_fin - costo_fin
    rent_fin = safe_div(margen_fin, produccion_fin)

    recaudado_cliente = float(fin.loc[fin["CLAS_CLIENTE"].eq("RECAUDADO"), "V.CLIENTE"].fillna(0).sum())
    pendiente_recaudar = float(fin.loc[fin["CLAS_CLIENTE"].eq("PENDIENTE POR RECAUDAR"), "V.CLIENTE"].fillna(0).sum())
    pct_recaudo = safe_div(recaudado_cliente, produccion_fin)

    pagado_terceros = float(fin.loc[fin["CLAS_PROVE"].eq("PAGADO A TERCEROS"), "V.CONDUCT"].fillna(0).sum())
    pendiente_pagar = float(fin.loc[fin["CLAS_PROVE"].eq("PENDIENTE POR PAGAR"), "V.CONDUCT"].fillna(0).sum())
    por_clasificar = float(fin.loc[fin["CLAS_PROVE"].eq("POR CLASIFICAR"), "V.CONDUCT"].fillna(0).sum())
    pct_pago_terceros = safe_div(pagado_terceros, costo_fin)
    pct_pendiente_pago = safe_div(pendiente_pagar, costo_fin)

    caja_operativa = recaudado_cliente - pagado_terceros
    brecha_pendiente = pendiente_recaudar - pendiente_pagar
    cobertura_obligaciones_caja = safe_div(caja_operativa, pendiente_pagar)

    tol_cop = 1.0
    tol_pct = 0.0001
    ctrl_margen = abs((produccion_fin - costo_fin) - margen_fin) < tol_cop
    ctrl_cliente = abs((recaudado_cliente + pendiente_recaudar) - produccion_fin) < tol_cop
    ctrl_prove = abs((pagado_terceros + pendiente_pagar + por_clasificar) - costo_fin) < tol_cop
    ctrl_rent = True if produccion_fin == 0 else abs((margen_fin / produccion_fin) - rent_fin) < tol_pct
    controles_ok = all([ctrl_margen, ctrl_cliente, ctrl_prove, ctrl_rent])

    # ---------------------------------------------------------
    # MENSAJE EJECUTIVO FINANCIERO
    # ---------------------------------------------------------
    fin_story_text = (
        f"La operación acumula {fmt_money(produccion_fin)} de producción con una "
        f"rentabilidad directa de {fmt_pct(rent_fin)}. "
        f"Se ha recaudado {fmt_money(recaudado_cliente)} ({fmt_pct(pct_recaudo)}) "
        f"y se han pagado {fmt_money(pagado_terceros)} a terceros. "
        f"La caja operativa del universo analizado es {fmt_money(caja_operativa)}; "
        f"quedan {fmt_money(pendiente_recaudar)} por recaudar y "
        f"{fmt_money(pendiente_pagar)} por pagar."
    )

    st.markdown(
        f"""
        <div class="story-banner">
            <div class="eyebrow">Lectura ejecutiva del periodo</div>
            <div class="story-text">{fin_story_text}</div>
        </div>
        <div class="story-path">
            <span class="node">Generamos</span><span class="arrow">→</span>
            <span class="node">Cobramos</span><span class="arrow">→</span>
            <span class="node">Pagamos</span><span class="arrow">→</span>
            <span class="node">Medimos riesgo</span><span class="arrow">→</span>
            <span class="node">Proyectamos</span>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.markdown("## 1. ¿Qué generamos? · Resultado económico")
    st.caption(
        "Primera lectura ejecutiva: cuánto se produjo, cuánto costó prestar los servicios "
        "y qué margen directo dejó la operación."
    )

    e1, e2, e3, e4 = st.columns(4)
    with e1:
        fin_card("Producción", fmt_money(produccion_fin), "Valor total de servicios · Σ V.CLIENTE")
    with e2:
        fin_card("Costo del servicio", fmt_money(costo_fin), "Costo directo asociado · Σ V.CONDUCT")
    with e3:
        fin_card("Margen directo", fmt_money(margen_fin), "Producción - Costo del Servicio", "good" if margen_fin >= 0 else "bad")
    with e4:
        fin_card("Rentabilidad directa", fmt_pct(rent_fin), "Margen directo / Producción", "good" if (pd.notna(rent_fin) and rent_fin >= 0) else "bad")

    st.markdown("## 2. ¿Cuánto convertimos en caja? · Recaudo y obligaciones")
    st.caption(
        "Segunda lectura: qué parte de la producción ya fue recaudada y qué parte del costo "
        "ya fue pagada a terceros."
    )

    st.markdown('<div class="subsection-tag">CLIENTES</div>', unsafe_allow_html=True)
    rc1, rc2, rc3 = st.columns(3)
    with rc1:
        fin_card("Recaudado cliente", fmt_money(recaudado_cliente), "ESTADO FA = FACT DEFINITIVA", "good")
    with rc2:
        fin_card("Pendiente por recaudar", fmt_money(pendiente_recaudar), f"{fmt_pct(safe_div(pendiente_recaudar, produccion_fin))} de la producción", "warn" if pendiente_recaudar > 0 else "good")
    with rc3:
        fin_card("% Recaudo", fmt_pct(pct_recaudo), "Recaudado / Producción", "good")

    st.markdown('<div class="subsection-tag">TERCEROS</div>', unsafe_allow_html=True)
    pt1, pt2, pt3 = st.columns(3)
    with pt1:
        fin_card("Pagado a terceros", fmt_money(pagado_terceros), "FAC PROVE = PAGADO / FACTURADO", "good")
    with pt2:
        fin_card("Pendiente por pagar", fmt_money(pendiente_pagar), f"{fmt_pct(pct_pendiente_pago)} del costo", "warn" if pendiente_pagar > 0 else "good")
    with pt3:
        fin_card("% Pago a terceros", fmt_pct(pct_pago_terceros), "Pagado a terceros / Costo del Servicio")

    st.markdown("## 3. ¿Qué nos queda? · Posición de caja operativa")
    cx1, cx2, cx3 = st.columns(3)
    with cx1:
        fin_card("Caja operativa", fmt_money(caja_operativa), "Recaudado cliente - Pagado a terceros", "good" if caja_operativa >= 0 else "bad")
    with cx2:
        fin_card("Cobertura de obligaciones", fmt_pct(cobertura_obligaciones_caja), "Caja operativa / Pendiente por Pagar", "good" if (pd.notna(cobertura_obligaciones_caja) and cobertura_obligaciones_caja >= 1) else "warn")
    with cx3:
        fin_card("Brecha pendiente", fmt_money(brecha_pendiente), "Pendiente por recaudar - Pendiente por pagar", "good" if brecha_pendiente >= 0 else "bad")

    st.markdown(
        '<div class="fin-panel"><strong>Lectura ejecutiva:</strong>'
        '<span class="fin-note"> la caja operativa corresponde al recaudo registrado menos pagos a terceros '
        'del universo filtrado. No equivale al saldo bancario total ni a la utilidad contable.</span></div>',
        unsafe_allow_html=True,
    )

    with st.expander("✓ Control matemático y calidad de estados", expanded=False):
        if controles_ok:
            st.success(
                "Las identidades principales cuadran dentro de la tolerancia definida: "
                "Producción = Recaudado + Pendiente; Costo = Pagado + Pendiente + Por clasificar; "
                "Margen = Producción - Costo."
            )
        else:
            st.error(
                "REVISAR DATOS: al menos una identidad financiera no cuadra dentro de la tolerancia definida."
            )

        if por_clasificar > tol_cop:
            unknown_states = sorted([
                x for x in fin.loc[
                    fin["CLAS_PROVE"].eq("POR CLASIFICAR"), "FAC_PROVE_N"
                ].dropna().unique().tolist() if x
            ])
            st.warning(
                "FAC PROVE contiene valores no clasificados. "
                f"Costo por clasificar: {fmt_money(por_clasificar)}. "
                "Estados detectados: "
                + (", ".join(unknown_states) if unknown_states else "VACÍO")
            )


    def grouped_financial(data: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
        base = data.groupby(group_cols, dropna=False).agg(Servicios=("CLIENTE", "size"), Produccion=("V.CLIENTE", "sum"), Costo=("V.CONDUCT", "sum")).reset_index()

        def add_group_sum(mask, value_col, out_col):
            s = data.loc[mask].groupby(group_cols, dropna=False)[value_col].sum().rename(out_col).reset_index()
            return s

        for mask, value_col, out_col in [
            (data["CLAS_CLIENTE"].eq("RECAUDADO"), "V.CLIENTE", "Recaudado"),
            (data["CLAS_CLIENTE"].eq("PENDIENTE POR RECAUDAR"), "V.CLIENTE", "Pendiente Recaudar"),
            (data["CLAS_PROVE"].eq("PAGADO A TERCEROS"), "V.CONDUCT", "Pagado Terceros"),
            (data["CLAS_PROVE"].eq("PENDIENTE POR PAGAR"), "V.CONDUCT", "Pendiente Pagar"),
            (data["CLAS_PROVE"].eq("POR CLASIFICAR"), "V.CONDUCT", "Por Clasificar"),
        ]:
            base = base.merge(add_group_sum(mask, value_col, out_col), on=group_cols, how="left")

        for c in ["Recaudado", "Pendiente Recaudar", "Pagado Terceros", "Pendiente Pagar", "Por Clasificar"]:
            base[c] = pd.to_numeric(base[c], errors="coerce").fillna(0)
        base["Margen"] = base["Produccion"] - base["Costo"]
        base["Rentabilidad"] = np.where(base["Produccion"].ne(0), base["Margen"] / base["Produccion"], np.nan)
        base["% Recaudo"] = np.where(base["Produccion"].ne(0), base["Recaudado"] / base["Produccion"], np.nan)
        base["% Pago Terceros"] = np.where(base["Costo"].ne(0), base["Pagado Terceros"] / base["Costo"], np.nan)
        return base

    cliente_fin = grouped_financial(fin, ["CLIENTE"]).sort_values("Produccion", ascending=False).reset_index(drop=True)

    st.markdown("## 4. ¿Cómo llegamos hasta aquí? · Evolución mensual")
    st.caption(
        "Comportamiento mes a mes usando FECHA CUM como fecha de corte financiero-operativo. "
        "FACTURADO corresponde a servicios en estado CUMPLIDO; RECAUDADO corresponde a ESTADO FA = FACT DEFINITIVA; "
        "PAGADO A TERCEROS corresponde a FAC PROVE = PAGADO/FACTURADO. "
        "Mientras no exista una fecha bancaria independiente para recaudo/pago proveedor, estos movimientos se atribuyen al mes de FECHA CUM."
    )

    evo = df_all.copy()
    evo = valid_services(evo)

    # Mantener los filtros de negocio, pero permitir observar todo el histórico mensual.
    evo_filters = [
        ("CLIENTE", f_client),
        ("Quien Creó", f_coord),
        ("ESTADO OP N", f_estado),
        ("Tipo Flota", f_flota),
        ("LINEA NEG", f_linea),
        ("T.NEGOCIO", f_negocio),
        ("TIPOLOGIA", f_tipologia),
        ("PLACA", f_placa),
        ("CONDUCTOR_NOMBRE", f_conductor),
    ]
    for _col, _vals in evo_filters:
        if _vals:
            evo = evo[evo[_col].astype(str).isin(_vals)]

    evo["ESTADO_FA_N"] = norm_fin_state(evo["ESTADO FA"])
    evo["FAC_PROVE_N"] = norm_fin_state(evo["FAC PROVE"])

    # FECHA CUM ya viene normalizada a datetime desde el cargue de datos.
    evo = evo[evo["FECHA CUM"].notna()].copy()
    evo["PERIODO_CUM"] = evo["FECHA CUM"].dt.to_period("M")

    # FACTURADO: valor cliente de servicios que alcanzaron CUMPLIDO.
    fact_m = (
        evo.loc[evo["ESTADO OP N"].eq("CUMPLIDO")]
        .groupby("PERIODO_CUM")["V.CLIENTE"].sum()
        .rename("Facturado")
    )

    # RECAUDADO: regla de negocio definida por Vigía.
    rec_m = (
        evo.loc[evo["ESTADO_FA_N"].eq("FACT DEFINITIVA")]
        .groupby("PERIODO_CUM")["V.CLIENTE"].sum()
        .rename("Recaudado")
    )

    # PAGADO A TERCEROS: estados PAGADO o FACTURADO en FAC PROVE.
    pay_m = (
        evo.loc[evo["FAC_PROVE_N"].isin(paid_supplier_states)]
        .groupby("PERIODO_CUM")["V.CONDUCT"].sum()
        .rename("Pagado Terceros")
    )

    srv_m = (
        evo.loc[evo["ESTADO OP N"].eq("CUMPLIDO")]
        .groupby("PERIODO_CUM").size()
        .rename("Servicios Cumplidos")
    )

    if len(evo):
        min_per = evo["PERIODO_CUM"].min()
        max_per = evo["PERIODO_CUM"].max()
        all_periods = pd.period_range(min_per, max_per, freq="M")
        monthly_evo = pd.DataFrame(index=all_periods)
        monthly_evo.index.name = "Periodo"
        monthly_evo = monthly_evo.join([fact_m, rec_m, pay_m, srv_m]).fillna(0).reset_index()
    else:
        monthly_evo = pd.DataFrame(columns=["Periodo","Facturado","Recaudado","Pagado Terceros","Servicios Cumplidos"])

    if not monthly_evo.empty:
        monthly_evo["Caja Operativa"] = monthly_evo["Recaudado"] - monthly_evo["Pagado Terceros"]
        monthly_evo["% Recaudo sobre Facturado"] = np.where(
            monthly_evo["Facturado"].ne(0),
            monthly_evo["Recaudado"] / monthly_evo["Facturado"],
            np.nan,
        )
        monthly_evo["Mes"] = monthly_evo["Periodo"].apply(
            lambda p: f"{MESES_ES[int(p.month)]} {int(p.year)}"
        )

        # Mostrar por defecto los últimos 12 meses disponibles.
        monthly_view = monthly_evo.tail(12).copy()

        fig_evo = go.Figure()
        fig_evo.add_trace(go.Bar(
            x=monthly_view["Mes"], y=monthly_view["Facturado"],
            name="Facturado", marker_color="#1769AA",
            hovertemplate="<b>%{x}</b><br>Facturado: $%{y:,.0f}<extra></extra>"
        ))
        fig_evo.add_trace(go.Bar(
            x=monthly_view["Mes"], y=monthly_view["Recaudado"],
            name="Recaudado", marker_color="#2DA66F",
            hovertemplate="<b>%{x}</b><br>Recaudado: $%{y:,.0f}<extra></extra>"
        ))
        fig_evo.add_trace(go.Bar(
            x=monthly_view["Mes"], y=monthly_view["Pagado Terceros"],
            name="Pagado a terceros", marker_color="#64B5F6",
            hovertemplate="<b>%{x}</b><br>Pagado a terceros: $%{y:,.0f}<extra></extra>"
        ))
        fig_evo.update_layout(
            barmode="group",
            height=450,
            margin=dict(l=20, r=20, t=50, b=90),
            paper_bgcolor="#07111F",
            plot_bgcolor="#07111F",
            font=dict(color="#E2E8F0"),
            legend=dict(orientation="h", y=1.12, x=0, font=dict(color="#F8FAFC", size=12)),
            xaxis=dict(title="", tickangle=-30, gridcolor="#1E3652"),
            yaxis=dict(title="Valor COP", gridcolor="#1E3652", tickformat="~s"),
            title="FACTURADO · RECAUDADO · PAGADO A TERCEROS — EVOLUCIÓN MENSUAL",
        )
        st.plotly_chart(fig_evo, use_container_width=True, key="monthly_financial_evolution")

        # Tarjetas del último mes disponible.
        last_row = monthly_evo.iloc[-1]
        m1, m2, m3, m4 = st.columns(4)
        with m1:
            fin_card("Facturado último mes", fmt_money(last_row["Facturado"]), str(last_row["Mes"]))
        with m2:
            fin_card("Recaudado último mes", fmt_money(last_row["Recaudado"]), f"{fmt_pct(last_row['% Recaudo sobre Facturado'])} del facturado", "good")
        with m3:
            fin_card("Pagado terceros último mes", fmt_money(last_row["Pagado Terceros"]), str(last_row["Mes"]))
        with m4:
            fin_card(
                "Caja operativa del mes",
                fmt_money(last_row["Caja Operativa"]),
                "Recaudado - Pagado a terceros",
                "good" if last_row["Caja Operativa"] >= 0 else "bad",
            )

        st.markdown('<div class="section-title">DETALLE MENSUAL</div>', unsafe_allow_html=True)
        evo_show = monthly_view[[
            "Mes","Servicios Cumplidos","Facturado","Recaudado","Pagado Terceros",
            "Caja Operativa","% Recaudo sobre Facturado"
        ]].copy()
        evo_show.columns = [
            "Mes","Servicios Cumplidos","Facturado","Recaudado","Pagado a Terceros",
            "Caja Operativa","% Recaudo"
        ]
        st.dataframe(
            evo_show.style.format({
                "Servicios Cumplidos": "{:,.0f}",
                "Facturado": lambda x: "$"+f"{x:,.0f}".replace(",", "."),
                "Recaudado": lambda x: "$"+f"{x:,.0f}".replace(",", "."),
                "Pagado a Terceros": lambda x: "$"+f"{x:,.0f}".replace(",", "."),
                "Caja Operativa": lambda x: "$"+f"{x:,.0f}".replace(",", "."),
                "% Recaudo": "{:.2%}",
            }),
            use_container_width=True,
            hide_index=True,
            height=min(470, 38 + 35 * max(1, len(evo_show))),
        )

        st.markdown(
            '<div class="fin-panel"><strong>Cómo leer esta gráfica:</strong>'
            '<span class="fin-note"> Facturado muestra el valor de los servicios que llegaron a CUMPLIDO en cada mes; '
            'Recaudado muestra cuánto de esos valores aparece como FACT DEFINITIVA; Pagado a Terceros muestra el costo '
            'asociado a registros marcados como PAGADO/FACTURADO. La diferencia Recaudado - Pagado a Terceros se muestra '
            'como Caja Operativa mensual del universo analizado.</span></div>',
            unsafe_allow_html=True,
        )
    else:
        st.info("No hay registros con FECHA CUM para construir la evolución financiera mensual.")


    # =========================================================
    # 5. ¿DÓNDE ESTÁ EL RIESGO?
    # =========================================================
    st.markdown("## 5. ¿Dónde está el riesgo? · Exposición prioritaria")
    st.caption(
        "Primero se identifican las concentraciones que requieren gestión: "
        "clientes con mayor saldo pendiente de recaudo y terceros con mayor obligación pendiente."
    )

    # Exposición de clientes.
    cliente_show = cliente_fin[
        [
            "CLIENTE", "Servicios", "Produccion", "Recaudado",
            "Pendiente Recaudar", "% Recaudo", "Costo",
            "Pagado Terceros", "Pendiente Pagar",
            "% Pago Terceros", "Margen", "Rentabilidad"
        ]
    ].copy()
    cliente_show.columns = [
        "Cliente", "Servicios", "Producción", "Recaudado",
        "Pendiente por Recaudar", "% Recaudo", "Costo del Servicio",
        "Pagado a Terceros", "Pendiente por Pagar",
        "% Pago Terceros", "Margen", "Rentabilidad"
    ]

    top_pending_client = (
        cliente_show.sort_values(
            ["Pendiente por Recaudar", "Producción"],
            ascending=[False, False],
        )
        .head(10)
        .copy()
    )

    # Exposición de terceros: vista base por conductor.
    supplier_by_driver = (
        grouped_financial(fin, ["CONDUCTOR_NOMBRE"])
        .sort_values(["Pendiente Pagar", "Costo"], ascending=[False, False])
        .reset_index(drop=True)
    )
    top_pending_sup = supplier_by_driver.head(10).copy()
    top_pending_sup["% Pendiente"] = np.where(
        top_pending_sup["Costo"].ne(0),
        top_pending_sup["Pendiente Pagar"] / top_pending_sup["Costo"],
        np.nan,
    )
    top_pending_sup = top_pending_sup[
        [
            "CONDUCTOR_NOMBRE", "Costo", "Pagado Terceros",
            "Pendiente Pagar", "% Pendiente"
        ]
    ].copy()
    top_pending_sup.columns = [
        "Conductor", "Costo del Servicio", "Pagado",
        "Pendiente por Pagar", "% Pendiente"
    ]

    rx1, rx2 = st.columns(2, gap="large")

    with rx1:
        st.markdown(
            '<div class="section-title">TOP CLIENTES · PENDIENTE POR RECAUDAR</div>',
            unsafe_allow_html=True,
        )
        st.dataframe(
            top_pending_client[
                [
                    "Cliente", "Producción", "Recaudado",
                    "Pendiente por Recaudar", "% Recaudo"
                ]
            ].style.format(
                {
                    "Producción": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Recaudado": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Pendiente por Recaudar": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "% Recaudo": "{:.2%}",
                }
            ),
            use_container_width=True,
            hide_index=True,
            height=390,
        )

    with rx2:
        st.markdown(
            '<div class="section-title">TOP TERCEROS · PENDIENTE POR PAGAR</div>',
            unsafe_allow_html=True,
        )
        st.dataframe(
            top_pending_sup.style.format(
                {
                    "Costo del Servicio": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Pagado": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Pendiente por Pagar": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "% Pendiente": "{:.2%}",
                }
            ),
            use_container_width=True,
            hide_index=True,
            height=390,
        )

    # Detalle de cartera y obligaciones queda disponible, pero no interrumpe la historia principal.
    with st.expander("🔎 Ver detalle de recaudo y pagos", expanded=False):
        st.markdown("### Recaudo por cliente")

        top_cli = cliente_fin.head(15)
        fig_cli = go.Figure()
        for col, name, color in [
            ("Produccion", "Producción", ROYAL_BLUE),
            ("Recaudado", "Recaudado", GREEN),
            ("Pendiente Recaudar", "Pendiente por Recaudar", AMBER),
        ]:
            fig_cli.add_trace(
                go.Bar(
                    x=top_cli["CLIENTE"],
                    y=top_cli[col],
                    name=name,
                    marker_color=color,
                    hovertemplate=f"<b>%{{x}}</b><br>{name}: $%{{y:,.0f}}<extra></extra>",
                )
            )
        fig_cli.update_layout(
            barmode="group",
            height=440,
            margin=dict(l=20, r=20, t=70, b=140),
            legend=dict(
                orientation="h",
                y=1.16,
                x=0,
                font=dict(color="#F8FAFC", size=12),
            ),
            xaxis=dict(tickangle=-45),
            yaxis=dict(title="COP"),
            bargap=.22,
            bargroupgap=.06,
        )
        st.plotly_chart(fig_cli, use_container_width=True, key="fin_cliente_chart")

        st.dataframe(
            cliente_show.style.format(
                {
                    "Servicios": "{:,.0f}",
                    "Producción": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Recaudado": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Pendiente por Recaudar": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Costo del Servicio": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Pagado a Terceros": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Pendiente por Pagar": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Margen": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "% Recaudo": "{:.2%}",
                    "% Pago Terceros": "{:.2%}",
                    "Rentabilidad": "{:.2%}",
                }
            ),
            use_container_width=True,
            hide_index=True,
            height=460,
        )

        st.markdown("### Pagos a terceros")
        supplier_mode = st.radio(
            "Analizar obligaciones por",
            ["CONDUCTOR", "PLACA", "CLIENTE"],
            horizontal=True,
            key="supplier_mode_fin",
        )
        supplier_col = {
            "CONDUCTOR": "CONDUCTOR_NOMBRE",
            "PLACA": "PLACA",
            "CLIENTE": "CLIENTE",
        }[supplier_mode]

        supplier_fin = (
            grouped_financial(fin, [supplier_col])
            .sort_values("Costo", ascending=False)
            .reset_index(drop=True)
        )
        top_sup = supplier_fin.head(15)

        fig_sup = go.Figure()
        for col, name, color in [
            ("Costo", "Costo del Servicio", ROYAL_BLUE),
            ("Pagado Terceros", "Pagado a Terceros", GREEN),
            ("Pendiente Pagar", "Pendiente por Pagar", AMBER),
        ]:
            fig_sup.add_trace(
                go.Bar(
                    x=top_sup[supplier_col],
                    y=top_sup[col],
                    name=name,
                    marker_color=color,
                    hovertemplate=f"<b>%{{x}}</b><br>{name}: $%{{y:,.0f}}<extra></extra>",
                )
            )
        fig_sup.update_layout(
            barmode="group",
            height=440,
            margin=dict(l=20, r=20, t=70, b=140),
            legend=dict(
                orientation="h",
                y=1.16,
                x=0,
                font=dict(color="#F8FAFC", size=12),
            ),
            xaxis=dict(tickangle=-45),
            yaxis=dict(title="COP"),
        )
        st.plotly_chart(
            fig_sup,
            use_container_width=True,
            key=f"fin_supplier_chart_{supplier_mode}",
        )

        tercero_det = (
            grouped_financial(fin, ["CONDUCTOR_NOMBRE", "PLACA", "CLIENTE"])
            .sort_values("Pendiente Pagar", ascending=False)
        )
        tercero_show = tercero_det[
            [
                "CONDUCTOR_NOMBRE", "PLACA", "CLIENTE", "Servicios",
                "Costo", "Pagado Terceros", "Pendiente Pagar",
                "% Pago Terceros"
            ]
        ].copy()
        tercero_show.columns = [
            "Conductor", "Placa", "Cliente", "Servicios",
            "Costo del Servicio", "Pagado a Terceros",
            "Pendiente por Pagar", "% Pago"
        ]

        st.dataframe(
            tercero_show.style.format(
                {
                    "Servicios": "{:,.0f}",
                    "Costo del Servicio": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Pagado a Terceros": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Pendiente por Pagar": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "% Pago": "{:.2%}",
                }
            ),
            use_container_width=True,
            hide_index=True,
            height=470,
        )

    # Para exportación, tercero_show debe existir aun si el expander no se abre.
    tercero_det = (
        grouped_financial(fin, ["CONDUCTOR_NOMBRE", "PLACA", "CLIENTE"])
        .sort_values("Pendiente Pagar", ascending=False)
    )
    tercero_show = tercero_det[
        [
            "CONDUCTOR_NOMBRE", "PLACA", "CLIENTE", "Servicios",
            "Costo", "Pagado Terceros", "Pendiente Pagar",
            "% Pago Terceros"
        ]
    ].copy()
    tercero_show.columns = [
        "Conductor", "Placa", "Cliente", "Servicios",
        "Costo del Servicio", "Pagado a Terceros",
        "Pendiente por Pagar", "% Pago"
    ]

    # =========================================================
    # 6. ¿QUÉ TAN RENTABLE ES CADA PARTE DEL NEGOCIO?
    # =========================================================
    st.markdown("## 6. ¿Qué tan rentable es cada parte del negocio?")
    st.caption(
        "Después de identificar la exposición, se revisa qué clientes y recursos "
        "explican el margen y la rentabilidad."
    )

    rent_client = (
        cliente_fin.sort_values("Produccion", ascending=False)
        .head(12)
        .copy()
    )

    pr1, pr2 = st.columns([1.05, 0.95], gap="large")

    with pr1:
        st.markdown(
            '<div class="section-title">MARGEN Y RENTABILIDAD · PRINCIPALES CLIENTES</div>',
            unsafe_allow_html=True,
        )

        fig_rent_cli = go.Figure()
        fig_rent_cli.add_trace(
            go.Bar(
                x=rent_client["CLIENTE"],
                y=rent_client["Margen"],
                name="Margen",
                marker_color=GREEN,
                hovertemplate="<b>%{x}</b><br>Margen: $%{y:,.0f}<extra></extra>",
            )
        )
        fig_rent_cli.add_trace(
            go.Scatter(
                x=rent_client["CLIENTE"],
                y=rent_client["Rentabilidad"] * 100,
                name="Rentabilidad %",
                mode="lines+markers",
                yaxis="y2",
                line=dict(color=SECONDARY_BLUE, width=3),
                marker=dict(size=8),
                hovertemplate="<b>%{x}</b><br>Rentabilidad: %{y:.2f}%<extra></extra>",
            )
        )
        fig_rent_cli.update_layout(
            height=440,
            margin=dict(l=20, r=55, t=70, b=140),
            legend=dict(
                orientation="h",
                y=1.16,
                x=0,
                font=dict(color="#F8FAFC", size=12),
            ),
            xaxis=dict(tickangle=-40),
            yaxis=dict(title="Margen COP"),
            yaxis2=dict(
                title="Rentabilidad %",
                overlaying="y",
                side="right",
                ticksuffix="%",
                showgrid=False,
            ),
        )
        st.plotly_chart(
            fig_rent_cli,
            use_container_width=True,
            key="fin_client_profitability",
        )

    with pr2:
        st.markdown(
            '<div class="section-title">DETALLE DE RENTABILIDAD POR CLIENTE</div>',
            unsafe_allow_html=True,
        )

        rent_client_show = rent_client[
            ["CLIENTE", "Produccion", "Costo", "Margen", "Rentabilidad"]
        ].copy()
        rent_client_show.columns = [
            "Cliente", "Producción", "Costo del Servicio",
            "Margen", "Rentabilidad"
        ]

        st.dataframe(
            rent_client_show.style.format(
                {
                    "Producción": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Costo del Servicio": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Margen": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Rentabilidad": "{:.2%}",
                }
            ),
            use_container_width=True,
            hide_index=True,
            height=440,
        )

    st.markdown("### Rentabilidad por placa y conductor")

    fin_resource_mode = st.radio(
        "Analizar resultado económico por",
        ["PLACA", "CONDUCTOR"],
        horizontal=True,
        key="fin_resource_economic_mode",
    )
    fin_resource_col = (
        "PLACA" if fin_resource_mode == "PLACA" else "CONDUCTOR_NOMBRE"
    )
    fin_resource_label = (
        "Placa" if fin_resource_mode == "PLACA" else "Conductor"
    )

    fin_resource = (
        grouped_financial(fin, [fin_resource_col])
        .sort_values(["Produccion", "Servicios"], ascending=[False, False])
        .reset_index(drop=True)
    )

    if not fin_resource.empty:
        top_resource = fin_resource.head(12).copy()

        fig_resource = go.Figure()
        for col, name, color in [
            ("Produccion", "Producción", ROYAL_BLUE),
            ("Costo", "Costo del Servicio", SECONDARY_BLUE),
            ("Margen", "Margen", GREEN),
        ]:
            fig_resource.add_trace(
                go.Bar(
                    x=top_resource[fin_resource_col],
                    y=top_resource[col],
                    name=name,
                    marker_color=color,
                    hovertemplate=f"<b>%{{x}}</b><br>{name}: $%{{y:,.0f}}<extra></extra>",
                )
            )

        fig_resource.update_layout(
            barmode="group",
            height=430,
            margin=dict(l=20, r=20, t=70, b=120),
            legend=dict(
                orientation="h",
                y=1.16,
                x=0,
                font=dict(color="#F8FAFC", size=12),
            ),
            xaxis=dict(tickangle=-40),
            yaxis=dict(title="COP"),
        )

        st.plotly_chart(
            fig_resource,
            use_container_width=True,
            key=f"fin_resource_econ_chart_{fin_resource_mode}",
        )

        fin_resource_show = fin_resource[
            [
                fin_resource_col, "Servicios", "Produccion", "Costo",
                "Margen", "Rentabilidad", "Recaudado", "% Recaudo",
                "Pagado Terceros", "% Pago Terceros"
            ]
        ].copy()
        fin_resource_show.columns = [
            fin_resource_label, "Servicios", "Producción",
            "Costo del Servicio", "Margen", "Rentabilidad",
            "Recaudado", "% Recaudo", "Pagado a Terceros",
            "% Pago Terceros"
        ]

        st.dataframe(
            fin_resource_show.style.format(
                {
                    "Servicios": "{:,.0f}",
                    "Producción": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Costo del Servicio": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Margen": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "Rentabilidad": "{:.2%}",
                    "Recaudado": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "% Recaudo": "{:.2%}",
                    "Pagado a Terceros": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
                    "% Pago Terceros": "{:.2%}",
                }
            ),
            use_container_width=True,
            hide_index=True,
            height=480,
        )

    # =========================================================
    # 7. ¿QUÉ ESPERAMOS? · PROYECCIÓN ESTADÍSTICA
    # =========================================================
    st.markdown("## 7. ¿Qué esperamos? · Proyección del próximo mes")
    st.caption("Proyección auditable construida con los últimos 6 meses completos. El escenario BASE usa tendencia lineal reciente; CONSERVADOR y OPTIMISTA reflejan la variabilidad histórica observada.")

    # ---------------------------------------------------------
    # Universo histórico para forecast
    # Se conservan filtros de negocio (cliente, coordinador, flota,
    # línea, tipo de negocio, tipología, placa y conductor), pero NO
    # se limita el histórico por Mes/Año/Rango de fecha seleccionado,
    # porque el modelo necesita meses anteriores completos.
    # ---------------------------------------------------------
    hist = df_all.copy()
    hist = valid_services(hist)
    hist_filters = [
        ("CLIENTE", f_client),
        ("Quien Creó", f_coord),
        ("ESTADO OP N", f_estado),
        ("Tipo Flota", f_flota),
        ("LINEA NEG", f_linea),
        ("T.NEGOCIO", f_negocio),
        ("TIPOLOGIA", f_tipologia),
        ("PLACA", f_placa),
        ("CONDUCTOR_NOMBRE", f_conductor),
    ]
    for _col, _vals in hist_filters:
        if _vals:
            hist = hist[hist[_col].astype(str).isin(_vals)]

    hist["ESTADO_FA_N"] = norm_fin_state(hist["ESTADO FA"])
    hist["FAC_PROVE_N"] = norm_fin_state(hist["FAC PROVE"])
    hist["CLAS_CLIENTE"] = np.where(hist["ESTADO_FA_N"].eq("FACT DEFINITIVA"), "RECAUDADO", "PENDIENTE POR RECAUDAR")
    hist["CLAS_PROVE"] = np.select(
        [hist["FAC_PROVE_N"].isin(paid_supplier_states), hist["FAC_PROVE_N"].eq("EN ESPERA")],
        ["PAGADO A TERCEROS", "PENDIENTE POR PAGAR"],
        default="POR CLASIFICAR",
    )

    today_fc = pd.Timestamp.now().normalize()
    current_month_start = today_fc.replace(day=1)
    last_complete_month_end = current_month_start - pd.Timedelta(days=1)
    forecast_month_start = current_month_start + pd.offsets.MonthBegin(1)
    forecast_month_name = f"{MESES_ES[int(forecast_month_start.month)]} {forecast_month_start.year}"

    # Últimos seis meses COMPLETOS anteriores al mes en curso.
    train_months = pd.period_range(end=last_complete_month_end.to_period("M"), periods=6, freq="M")
    hist_train = hist[hist["FECHA_CARGA"].dt.to_period("M").isin(train_months)].copy()
    hist_train["PERIODO_M"] = hist_train["FECHA_CARGA"].dt.to_period("M")

    def monthly_forecast_frame(data: pd.DataFrame) -> pd.DataFrame:
        rows = []
        for per in train_months:
            m = data[data["PERIODO_M"].eq(per)] if "PERIODO_M" in data.columns else data[data["FECHA_CARGA"].dt.to_period("M").eq(per)]
            prod_m = float(m["V.CLIENTE"].fillna(0).sum())
            cost_m = float(m["V.CONDUCT"].fillna(0).sum())
            rec_m = float(m.loc[m["CLAS_CLIENTE"].eq("RECAUDADO"), "V.CLIENTE"].fillna(0).sum()) if len(m) else 0.0
            pay_m = float(m.loc[m["CLAS_PROVE"].eq("PAGADO A TERCEROS"), "V.CONDUCT"].fillna(0).sum()) if len(m) else 0.0
            rows.append({"Periodo": per, "Producción": prod_m, "Costo": cost_m, "Recaudado": rec_m, "Pagado": pay_m, "Servicios": len(m)})
        out = pd.DataFrame(rows)
        out["Ratio Costo"] = np.where(out["Producción"].ne(0), out["Costo"] / out["Producción"], np.nan)
        out["Ratio Recaudo"] = np.where(out["Producción"].ne(0), out["Recaudado"] / out["Producción"], np.nan)
        out["Ratio Pago"] = np.where(out["Costo"].ne(0), out["Pagado"] / out["Costo"], np.nan)
        return out

    monthly_fc = monthly_forecast_frame(hist_train)
    nonzero_prod = monthly_fc[monthly_fc["Producción"].gt(0)].copy()

    def weighted_ratio(series: pd.Series, default=np.nan):
        s = pd.to_numeric(series, errors="coerce").dropna()
        if s.empty:
            return default
        w = np.arange(1, len(s) + 1, dtype=float)
        return float(np.average(s.to_numpy(dtype=float), weights=w))

    def trend_forecast(values: pd.Series):
        y = pd.to_numeric(values, errors="coerce").fillna(0).to_numpy(dtype=float)
        n = len(y)
        if n == 0:
            return 0.0, 0.0, "BAJA", 0.0
        if n == 1:
            return max(0.0, y[-1]), 0.0, "BAJA", 0.0
        x = np.arange(n, dtype=float)
        slope, intercept = np.polyfit(x, y, 1)
        fitted = intercept + slope * x
        base = max(0.0, float(intercept + slope * n))
        rmse = float(np.sqrt(np.mean((y - fitted) ** 2)))
        pct = pd.Series(y).replace(0, np.nan).pct_change(fill_method=None).replace([np.inf, -np.inf], np.nan).dropna()
        vol_amount = abs(base) * float(pct.std(ddof=1)) if len(pct) >= 2 and pd.notna(pct.std(ddof=1)) else 0.0
        uncertainty = max(rmse, vol_amount)
        rel = safe_div(uncertainty, base)
        if pd.isna(rel):
            confidence = "BAJA"
        elif rel <= 0.10:
            confidence = "ALTA"
        elif rel <= 0.20:
            confidence = "MEDIA"
        else:
            confidence = "BAJA"
        return base, uncertainty, confidence, float(slope)

    if len(nonzero_prod) < 3:
        st.warning("Histórico insuficiente para una proyección estadística robusta: se requieren al menos 3 meses completos con producción. El módulo mostrará la mejor estimación disponible y marcará confianza BAJA.")

    base_prod_fc, uncertainty_fc, confidence_fc, slope_fc = trend_forecast(monthly_fc["Producción"])
    cons_prod_fc = max(0.0, base_prod_fc - uncertainty_fc)
    opt_prod_fc = max(0.0, base_prod_fc + uncertainty_fc)

    ratio_cost_fc = weighted_ratio(monthly_fc["Ratio Costo"], safe_div(costo_fin, produccion_fin))
    ratio_rec_fc = weighted_ratio(monthly_fc["Ratio Recaudo"], safe_div(recaudado_cliente, produccion_fin))
    ratio_pay_fc = weighted_ratio(monthly_fc["Ratio Pago"], safe_div(pagado_terceros, costo_fin))
    ratio_cost_fc = float(np.clip(ratio_cost_fc if pd.notna(ratio_cost_fc) else 0, 0, 5))
    ratio_rec_fc = float(np.clip(ratio_rec_fc if pd.notna(ratio_rec_fc) else 0, 0, 1))
    ratio_pay_fc = float(np.clip(ratio_pay_fc if pd.notna(ratio_pay_fc) else 0, 0, 1))

    scenario = st.radio("Escenario", ["CONSERVADOR", "BASE", "OPTIMISTA"], index=1, horizontal=True, key="forecast_scenario")
    prod_by_scenario = {"CONSERVADOR": cons_prod_fc, "BASE": base_prod_fc, "OPTIMISTA": opt_prod_fc}
    proj_prod = prod_by_scenario[scenario]
    proj_recaudo = proj_prod * ratio_rec_fc
    proj_cost = proj_prod * ratio_cost_fc
    proj_paid = proj_cost * ratio_pay_fc
    proj_margin = proj_prod - proj_cost
    proj_rent = safe_div(proj_margin, proj_prod)
    proj_cash = proj_recaudo - proj_paid
    proj_pending_rec = proj_prod - proj_recaudo
    proj_pending_pay = proj_cost - proj_paid

    conf_tone = "good" if confidence_fc == "ALTA" else ("warn" if confidence_fc == "MEDIA" else "bad")
    st.markdown(f'<div class="fin-panel"><strong>Mes proyectado:</strong> {forecast_month_name} &nbsp;·&nbsp; <strong>Confianza del modelo:</strong> {confidence_fc} &nbsp;·&nbsp; <span class="fin-note">La confianza se determina por la variabilidad/residuo del histórico de producción. Los escenarios cambian el volumen de producción; los ratios financieros usan promedios históricos ponderados para evitar supuestos arbitrarios.</span></div>', unsafe_allow_html=True)

    f1, f2, f3, f4 = st.columns(4)
    with f1: fin_card("Producción proyectada", fmt_money(proj_prod), f"Escenario {scenario}", conf_tone)
    with f2: fin_card("Recaudo proyectado", fmt_money(proj_recaudo), f"Proporción histórica esperada: {fmt_pct(ratio_rec_fc)}", "good")
    with f3: fin_card("Costo proyectado", fmt_money(proj_cost), f"Proporción histórica costo/producción: {fmt_pct(ratio_cost_fc)}")
    with f4: fin_card("Pago proyectado terceros", fmt_money(proj_paid), f"Proporción histórica de pago: {fmt_pct(ratio_pay_fc)}")
    f5, f6, f7, f8 = st.columns(4)
    with f5: fin_card("Margen proyectado", fmt_money(proj_margin), "Producción proyectada - Costo proyectado", "good" if proj_margin >= 0 else "bad")
    with f6: fin_card("Rentabilidad proyectada", fmt_pct(proj_rent), "Margen proyectado / Producción proyectada", "good" if pd.notna(proj_rent) and proj_rent >= 0 else "bad")
    with f7: fin_card("Caja operativa proyectada", fmt_money(proj_cash), "Recaudo proyectado - Pago proyectado a terceros", "good" if proj_cash >= 0 else "bad")
    with f8: fin_card("Pendiente por recaudar proyectado", fmt_money(proj_pending_rec), "Producción proyectada - Recaudo proyectado", "warn" if proj_pending_rec > 0 else "good")

    # Comparativo contra el último mes completo real.
    last_row = monthly_fc.iloc[-1]
    last_label = str(last_row["Periodo"])
    comparison = pd.DataFrame([
        ["Producción", float(last_row["Producción"]), proj_prod],
        ["Recaudo", float(last_row["Recaudado"]), proj_recaudo],
        ["Costo del Servicio", float(last_row["Costo"]), proj_cost],
        ["Pago a Terceros", float(last_row["Pagado"]), proj_paid],
        ["Margen", float(last_row["Producción"] - last_row["Costo"]), proj_margin],
        ["Caja Operativa", float(last_row["Recaudado"] - last_row["Pagado"]), proj_cash],
    ], columns=["Indicador", f"Último mes completo ({last_label})", f"Proyección {forecast_month_name}"])
    comparison["Variación %"] = np.where(
        comparison[f"Último mes completo ({last_label})"].ne(0),
        comparison[f"Proyección {forecast_month_name}"] / comparison[f"Último mes completo ({last_label})"] - 1,
        np.nan,
    )
    last_rent = safe_div(float(last_row["Producción"] - last_row["Costo"]), float(last_row["Producción"]))
    comparison_rent = pd.DataFrame({
        "Indicador": ["Rentabilidad"],
        f"Último mes completo ({last_label})": [last_rent],
        f"Proyección {forecast_month_name}": [proj_rent],
        "Variación p.p.": [(proj_rent - last_rent) * 100 if pd.notna(proj_rent) and pd.notna(last_rent) else np.nan],
    })

    st.markdown('<div class="section-title">ÚLTIMO MES COMPLETO VS PRÓXIMO MES PROYECTADO</div>', unsafe_allow_html=True)
    st.dataframe(
        comparison.style.format({
            f"Último mes completo ({last_label})": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
            f"Proyección {forecast_month_name}": lambda x: "$" + f"{x:,.0f}".replace(",", "."),
            "Variación %": "{:+.2%}",
        }),
        use_container_width=True, hide_index=True,
    )
    st.dataframe(
        comparison_rent.style.format({
            f"Último mes completo ({last_label})": "{:.2%}",
            f"Proyección {forecast_month_name}": "{:.2%}",
            "Variación p.p.": "{:+.2f}",
        }),
        use_container_width=True, hide_index=True,
    )

    # Histórico de 6 meses + forecast BASE para hacer visible la tendencia.
    fig_fc = go.Figure()
    x_hist = [str(p) for p in monthly_fc["Periodo"]]
    fig_fc.add_trace(go.Scatter(x=x_hist, y=monthly_fc["Producción"], mode="lines+markers", name="Producción histórica", line=dict(color="#64B5F6", width=3), marker=dict(size=8)))
    fig_fc.add_trace(go.Scatter(x=[x_hist[-1], forecast_month_name], y=[float(last_row["Producción"]), base_prod_fc], mode="lines+markers", name="Proyección BASE", line=dict(color="#2DA66F", width=3, dash="dash"), marker=dict(size=10)))
    fig_fc.add_trace(go.Scatter(x=[forecast_month_name, forecast_month_name], y=[cons_prod_fc, opt_prod_fc], mode="lines+markers", name="Rango escenarios", line=dict(color="#D98E04", width=7), marker=dict(size=10), hovertemplate="Rango: $%{y:,.0f}<extra></extra>"))
    fig_fc.update_layout(
        height=420,
        margin=dict(l=20, r=20, t=78, b=45),
        paper_bgcolor="#07111F",
        plot_bgcolor="#101C2C",
        font=dict(color="#F8FAFC"),
        legend=dict(
            orientation="h",
            y=1.18,
            x=0,
            font=dict(color="#F8FAFC", size=12),
            bgcolor="rgba(7,17,31,0.94)",
            bordercolor="#29415E",
            borderwidth=1,
        ),
        yaxis=dict(
            title="Producción COP",
            gridcolor="#29415E",
            tickfont=dict(color="#CBD5E1"),
            title_font=dict(color="#CBD5E1"),
        ),
        xaxis=dict(
            gridcolor="#29415E",
            tickfont=dict(color="#CBD5E1"),
        ),
    )
    st.plotly_chart(fig_fc, use_container_width=True, key="forecast_production_history")

    # Proyección por cliente: tendencia individual de producción + ratios históricos propios.
    st.markdown('<div class="section-title">PROYECCIÓN BASE POR CLIENTE — ¿QUIÉN EXPLICA EL PRÓXIMO MES?</div>', unsafe_allow_html=True)
    client_rows = []
    for cli in sorted(hist_train["CLIENTE"].dropna().astype(str).unique().tolist()):
        cd = hist_train[hist_train["CLIENTE"].astype(str).eq(cli)].copy()
        cm = monthly_forecast_frame(cd)
        cbase, _, cconf, _ = trend_forecast(cm["Producción"])
        cr_cost = weighted_ratio(cm["Ratio Costo"], ratio_cost_fc)
        cr_rec = weighted_ratio(cm["Ratio Recaudo"], ratio_rec_fc)
        cr_cost = float(np.clip(cr_cost if pd.notna(cr_cost) else ratio_cost_fc, 0, 5))
        cr_rec = float(np.clip(cr_rec if pd.notna(cr_rec) else ratio_rec_fc, 0, 1))
        ccost = cbase * cr_cost
        crec = cbase * cr_rec
        cmarg = cbase - ccost
        client_rows.append({"Cliente":cli, "Producción Proyectada":cbase, "Recaudo Proyectado":crec, "Costo Proyectado":ccost, "Margen Proyectado":cmarg, "Rentabilidad":safe_div(cmarg,cbase), "Confianza":cconf})
    client_forecast = pd.DataFrame(client_rows)
    if not client_forecast.empty:
        total_cli_proj = float(client_forecast["Producción Proyectada"].sum())
        client_forecast["Participación Proyectada %"] = np.where(total_cli_proj != 0, client_forecast["Producción Proyectada"] / total_cli_proj, np.nan)
        client_forecast = client_forecast.sort_values("Producción Proyectada", ascending=False).reset_index(drop=True)
        st.dataframe(client_forecast.head(20).style.format({
            "Producción Proyectada":lambda x:"$"+f"{x:,.0f}".replace(",","."),
            "Recaudo Proyectado":lambda x:"$"+f"{x:,.0f}".replace(",","."),
            "Costo Proyectado":lambda x:"$"+f"{x:,.0f}".replace(",","."),
            "Margen Proyectado":lambda x:"$"+f"{x:,.0f}".replace(",","."),
            "Rentabilidad":"{:.2%}", "Participación Proyectada %":"{:.2%}"
        }), use_container_width=True, hide_index=True, height=520)
    else:
        client_forecast = pd.DataFrame(columns=["Cliente","Producción Proyectada","Recaudo Proyectado","Costo Proyectado","Margen Proyectado","Rentabilidad","Confianza","Participación Proyectada %"])
        st.info("No hay histórico suficiente por cliente para construir el detalle de proyección.")

    with st.expander("🔎 Ver metodología de proyección", expanded=False):
        st.markdown(f"""
**Horizonte:** {forecast_month_name}.  
**Histórico utilizado:** últimos 6 meses completos anteriores al mes en curso: {', '.join(str(p) for p in train_months)}.  
**Producción BASE:** tendencia lineal por mínimos cuadrados sobre la producción mensual.  
**Rango CONSERVADOR/OPTIMISTA:** BASE ± incertidumbre observada; la incertidumbre toma el mayor valor entre RMSE de la tendencia y volatilidad mensual reciente aplicada al forecast. El escenario conservador nunca baja de cero.  
**Costo:** Producción proyectada × ratio histórico ponderado Costo/Producción.  
**Recaudo:** Producción proyectada × ratio histórico ponderado Recaudado/Producción.  
**Pago a terceros:** Costo proyectado × ratio histórico ponderado Pagado/Costo.  
**Margen:** Producción − Costo.  
**Rentabilidad:** Margen / Producción.  
**Caja operativa proyectada:** Recaudo proyectado − Pago proyectado a terceros.  
**Ponderación de proporciones históricas:** los meses más recientes reciben mayor peso (1,2,3,4,5,6).  
**Confianza:** ALTA si la incertidumbre relativa ≤10%; MEDIA si ≤20%; BAJA si >20% o el histórico es insuficiente.

**Advertencia metodológica:** sin una columna de fecha efectiva de recaudo/pago, el proyección de recaudo y pago estima la proporción esperada a partir de los estados financieros observados en la base; no modela todavía el desfase exacto en días de entrada o salida de caja.
        """)
        methodology_table = monthly_fc.copy()
        methodology_table["Periodo"] = methodology_table["Periodo"].astype(str)
        st.dataframe(methodology_table.style.format({"Producción":lambda x:"$"+f"{x:,.0f}".replace(",","."), "Costo":lambda x:"$"+f"{x:,.0f}".replace(",","."), "Recaudado":lambda x:"$"+f"{x:,.0f}".replace(",","."), "Pagado":lambda x:"$"+f"{x:,.0f}".replace(",","."), "Ratio Costo":"{:.2%}", "Ratio Recaudo":"{:.2%}", "Ratio Pago":"{:.2%}"}), use_container_width=True, hide_index=True)

    projection_table = pd.DataFrame([
        {"Escenario":"CONSERVADOR", "Producción":cons_prod_fc, "Recaudo":cons_prod_fc*ratio_rec_fc, "Costo del Servicio":cons_prod_fc*ratio_cost_fc, "Pago a Terceros":cons_prod_fc*ratio_cost_fc*ratio_pay_fc, "Margen":cons_prod_fc-(cons_prod_fc*ratio_cost_fc), "Rentabilidad":safe_div(cons_prod_fc-(cons_prod_fc*ratio_cost_fc),cons_prod_fc), "Caja Operativa":cons_prod_fc*ratio_rec_fc-(cons_prod_fc*ratio_cost_fc*ratio_pay_fc)},
        {"Escenario":"BASE", "Producción":base_prod_fc, "Recaudo":base_prod_fc*ratio_rec_fc, "Costo del Servicio":base_prod_fc*ratio_cost_fc, "Pago a Terceros":base_prod_fc*ratio_cost_fc*ratio_pay_fc, "Margen":base_prod_fc-(base_prod_fc*ratio_cost_fc), "Rentabilidad":safe_div(base_prod_fc-(base_prod_fc*ratio_cost_fc),base_prod_fc), "Caja Operativa":base_prod_fc*ratio_rec_fc-(base_prod_fc*ratio_cost_fc*ratio_pay_fc)},
        {"Escenario":"OPTIMISTA", "Producción":opt_prod_fc, "Recaudo":opt_prod_fc*ratio_rec_fc, "Costo del Servicio":opt_prod_fc*ratio_cost_fc, "Pago a Terceros":opt_prod_fc*ratio_cost_fc*ratio_pay_fc, "Margen":opt_prod_fc-(opt_prod_fc*ratio_cost_fc), "Rentabilidad":safe_div(opt_prod_fc-(opt_prod_fc*ratio_cost_fc),opt_prod_fc), "Caja Operativa":opt_prod_fc*ratio_rec_fc-(opt_prod_fc*ratio_cost_fc*ratio_pay_fc)},
    ])

    st.info("La proyección estadística es una estimación basada en comportamiento histórico; no constituye producción, recaudo ni pagos garantizados. El escenario BASE es la referencia central y los escenarios CONSERVADOR/OPTIMISTA expresan el rango derivado de la variabilidad observada.")

    st.markdown("## 8. Detalle y descarga ejecutiva")
    resumen_export = pd.DataFrame([
        ["Servicios", servicios_fin], ["Producción", produccion_fin], ["Recaudado cliente", recaudado_cliente], ["Pendiente por recaudar", pendiente_recaudar], ["% Recaudo", pct_recaudo], ["Costo del Servicio", costo_fin], ["Pagado a terceros", pagado_terceros], ["Pendiente por pagar", pendiente_pagar], ["Por clasificar", por_clasificar], ["% Pago a terceros", pct_pago_terceros], ["Margen", margen_fin], ["Rentabilidad", rent_fin], ["Caja operativa", caja_operativa], ["Brecha pendiente", brecha_pendiente], ["Cobertura de obligaciones con caja", cobertura_obligaciones_caja]
    ], columns=["Indicador","Valor"])

    fin_excel = BytesIO()
    with pd.ExcelWriter(fin_excel, engine="openpyxl") as writer:
        resumen_export.to_excel(writer, index=False, sheet_name="Resumen")
        cliente_show.to_excel(writer, index=False, sheet_name="Clientes")
        tercero_show.to_excel(writer, index=False, sheet_name="Terceros")
        projection_table.to_excel(writer, index=False, sheet_name="Forecast_Escenarios")
        comparison.to_excel(writer, index=False, sheet_name="Forecast_Comparativo")
        comparison_rent.to_excel(writer, index=False, sheet_name="Forecast_Rentabilidad")
        client_forecast.to_excel(writer, index=False, sheet_name="Forecast_Clientes")
        methodology_table.to_excel(writer, index=False, sheet_name="Forecast_Metodologia")
        pd.DataFrame({"Control":["Margen", "Cliente", "Terceros", "Rentabilidad"], "Resultado":[ctrl_margen, ctrl_cliente, ctrl_prove, ctrl_rent]}).to_excel(writer, index=False, sheet_name="Auditoria")
        for ws in writer.book.worksheets:
            fill = openpyxl.styles.PatternFill(fill_type="solid", fgColor="003B8E")
            font = openpyxl.styles.Font(color="FFFFFF", bold=True)
            for cell in ws[1]:
                cell.fill = fill
                cell.font = font
                cell.alignment = openpyxl.styles.Alignment(horizontal="center")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col_cells in ws.columns:
                max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
                ws.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 12), 42)
    fin_excel.seek(0)

    st.download_button("⬇️ Descargar Posición Financiera y Proyecciones", data=fin_excel, file_name="Posicion_Financiera_y_Proyecciones_VSE.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
    st.stop()


# =========================================================
# CENTRO DE CONTROL · STORYTELLING OPERATIVO
# =========================================================

prod = len(valid)

cumplidos = int(valid["ESTADO OP N"].eq("CUMPLIDO").sum())
cumpl_oper = int(valid["ESTADO OP N"].eq("CUMPLIDO OPERATIVO").sum())
en_programacion = int(valid["ESTADO OP N"].eq("EN PROGRAMACION").sum())
en_transito = int(valid["ESTADO OP N"].eq("EN TRANSITO").sum())
pend_oper = en_programacion + en_transito

cierre_final = safe_div(cumplidos, prod)

if "F PRECUMP" in valid.columns:
    ico_precump_count = int(valid["F PRECUMP"].notna().sum())
else:
    ico_precump_count = 0
ico_precump = safe_div(ico_precump_count, prod)

trend = aggregate_period(valid, grain)
last = trend.iloc[-1]
prev = trend.iloc[-2] if len(trend) > 1 else None
d_servicios = last["Var_Servicios"] if prev is not None else np.nan

ultimo_periodo_incompleto = bool(
    "Periodo_Incompleto" in trend.columns
    and len(trend)
    and trend.iloc[-1]["Periodo_Incompleto"]
)


def operational_period_frame(data: pd.DataFrame, period_grain: str) -> pd.DataFrame:
    x = data.copy()

    if period_grain == "Mensual":
        x["PERIODO_ORDEN"] = x["AÑO"] * 100 + x["MES_NUM"]
        x["PERIODO"] = x.apply(
            lambda r: f"{MONTHS_ES.get(r['MES_NUM'], '')} {int(r['AÑO'])}",
            axis=1,
        )
    elif period_grain == "Bimestral":
        x["PERIODO_ORDEN"] = x["AÑO"] * 10 + x["BIMESTRE_NUM"].astype(int)
        x["PERIODO"] = x.apply(
            lambda r: f"B{int(r['BIMESTRE_NUM'])} {int(r['AÑO'])}",
            axis=1,
        )
    elif period_grain == "Trimestral":
        x["PERIODO_ORDEN"] = x["AÑO"] * 10 + x["TRIMESTRE_NUM"].astype(int)
        x["PERIODO"] = x.apply(
            lambda r: f"T{int(r['TRIMESTRE_NUM'])} {int(r['AÑO'])}",
            axis=1,
        )
    else:
        x["PERIODO_ORDEN"] = x["AÑO"] * 10 + x["SEMESTRE_NUM"].astype(int)
        x["PERIODO"] = x.apply(
            lambda r: f"S{int(r['SEMESTRE_NUM'])} {int(r['AÑO'])}",
            axis=1,
        )

    x["_CUMPLIDO"] = x["ESTADO OP N"].eq("CUMPLIDO").astype(int)
    x["_CUMPLIDO_OPER"] = x["ESTADO OP N"].eq("CUMPLIDO OPERATIVO").astype(int)
    x["_PROGRAMACION"] = x["ESTADO OP N"].eq("EN PROGRAMACION").astype(int)
    x["_TRANSITO"] = x["ESTADO OP N"].eq("EN TRANSITO").astype(int)
    x["_PENDIENTE"] = x["_PROGRAMACION"] + x["_TRANSITO"]
    x["_ICO"] = (
        x["F PRECUMP"].notna().astype(int)
        if "F PRECUMP" in x.columns
        else 0
    )

    g = (
        x.groupby(["PERIODO_ORDEN", "PERIODO"], dropna=False)
        .agg(
            Servicios=("CLIENTE", "size"),
            Cumplido=("_CUMPLIDO", "sum"),
            Cumplido_Operativo=("_CUMPLIDO_OPER", "sum"),
            En_Programacion=("_PROGRAMACION", "sum"),
            En_Transito=("_TRANSITO", "sum"),
            Pendientes=("_PENDIENTE", "sum"),
            ICO_Count=("_ICO", "sum"),
        )
        .reset_index()
        .sort_values("PERIODO_ORDEN")
    )

    g["Cierre Final %"] = np.where(
        g["Servicios"].ne(0), g["Cumplido"] / g["Servicios"], np.nan
    )
    g["Pendientes %"] = np.where(
        g["Servicios"].ne(0), g["Pendientes"] / g["Servicios"], np.nan
    )
    g["ICO a Corte %"] = np.where(
        g["Servicios"].ne(0), g["ICO_Count"] / g["Servicios"], np.nan
    )
    return g


op_period = operational_period_frame(valid, grain)


def performance_by_group(data: pd.DataFrame, group_col: str) -> pd.DataFrame:
    x = data.copy()
    x["_CUMPLIDO"] = x["ESTADO OP N"].eq("CUMPLIDO").astype(int)
    x["_CUMPLIDO_OPER"] = x["ESTADO OP N"].eq("CUMPLIDO OPERATIVO").astype(int)
    x["_PENDIENTE"] = x["ESTADO OP N"].isin(
        ["EN PROGRAMACION", "EN TRANSITO"]
    ).astype(int)

    g = (
        x.groupby(group_col, dropna=False)
        .agg(
            Servicios=("CLIENTE", "size"),
            Cumplidos=("_CUMPLIDO", "sum"),
            Cumplido_Operativo=("_CUMPLIDO_OPER", "sum"),
            Pendientes=("_PENDIENTE", "sum"),
        )
        .reset_index()
    )

    g["Cierre %"] = np.where(
        g["Servicios"].ne(0),
        g["Cumplidos"] / g["Servicios"],
        np.nan,
    )
    g["Pendientes %"] = np.where(
        g["Servicios"].ne(0),
        g["Pendientes"] / g["Servicios"],
        np.nan,
    )
    g["Cumplido Operativo %"] = np.where(
        g["Servicios"].ne(0),
        g["Cumplido_Operativo"] / g["Servicios"],
        np.nan,
    )
    return g


client_perf = performance_by_group(valid, "CLIENTE")
coord_perf = performance_by_group(valid, "Quien Creó")


# ---------------------------------------------------------
# Mensaje ejecutivo automático
# ---------------------------------------------------------
top_pending_client_row = (
    client_perf.sort_values(
        ["Pendientes", "Servicios"],
        ascending=[False, False],
    ).iloc[0]
    if not client_perf.empty
    else None
)

top_pending_coord_row = (
    coord_perf.sort_values(
        ["Pendientes", "Servicios"],
        ascending=[False, False],
    ).iloc[0]
    if not coord_perf.empty
    else None
)

client_focus_name = (
    str(top_pending_client_row["CLIENTE"])
    if top_pending_client_row is not None
    else "Sin dato"
)
client_focus_pending = (
    int(top_pending_client_row["Pendientes"])
    if top_pending_client_row is not None
    else 0
)

op_story_text = (
    f"Se registran {fmt_int(prod)} servicios válidos, con cierre final de "
    f"{fmt_pct(cierre_final)}. Los pendientes operativos son {fmt_int(pend_oper)} "
    f"({fmt_pct(safe_div(pend_oper, prod))}) y el ICO a corte es "
    f"{fmt_pct(ico_precump)}. La mayor concentración de pendientes está en "
    f"{client_focus_name} con {fmt_int(client_focus_pending)} servicios."
)

st.markdown(
    f"""
    <div class="story-banner">
        <div class="eyebrow">Lectura ejecutiva del periodo</div>
        <div class="story-text">{op_story_text}</div>
    </div>
    <div class="story-path">
        <span class="node">Qué pasó</span><span class="arrow">→</span>
        <span class="node">Cómo cerramos</span><span class="arrow">→</span>
        <span class="node">Dónde está el problema</span><span class="arrow">→</span>
        <span class="node">Quién lo explica</span><span class="arrow">→</span>
        <span class="node">Dónde actuar</span>
    </div>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# 1. ¿CÓMO ESTAMOS?
# =========================================================
st.markdown("## 1. ¿Cómo estamos? · Resumen operativo")
st.caption(
    "Cuatro indicadores para entender el estado general antes de entrar al detalle."
)

k1, k2, k3, k4 = st.columns(4)

with k1:
    kpi_card(
        "Total Servicios",
        fmt_int(prod),
        d_servicios,
        True,
        "Servicios válidos · excluye ANULADO",
    )
with k2:
    kpi_card(
        "Cierre Final",
        fmt_pct(cierre_final),
        note="CUMPLIDO / Servicios válidos",
        show_delta=False,
    )
with k3:
    kpi_card(
        "Pendientes Operativos",
        fmt_int(pend_oper),
        note=f"{fmt_pct(safe_div(pend_oper, prod))} · Programación + Tránsito",
        show_delta=False,
    )
with k4:
    kpi_card(
        "ICO a Corte",
        fmt_pct(ico_precump),
        note="F PRECUMP / Servicios válidos",
        show_delta=False,
    )

if ultimo_periodo_incompleto:
    st.info(
        "📌 El último periodo está en curso. La variación de servicios no se compara "
        "contra un periodo completo hasta el cierre."
    )


# =========================================================
# 2. PRIMERA DIAGONAL · ¿QUÉ PASÓ? / ¿CÓMO CERRAMOS?
# =========================================================
st.markdown("## 2. ¿Qué pasó y cómo cerramos?")
st.caption(
    "La lectura empieza por el volumen y continúa hacia la calidad del cierre."
)

z1, z2 = st.columns([1.08, 0.92], gap="large")

with z1:
    st.markdown('<div class="question-tag">¿Qué pasó?</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-title">SERVICIOS POR PERIODO + VARIACIÓN</div>',
        unsafe_allow_html=True,
    )
    st.plotly_chart(
        trend_chart(
            trend,
            "Servicios",
            "Var_Servicios",
            "Servicios",
            money=False,
        ),
        use_container_width=True,
        key="story_services_period",
    )

with z2:
    st.markdown('<div class="question-tag">¿Cómo cerramos?</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-title">CIERRE · ICO · PENDIENTES POR PERIODO</div>',
        unsafe_allow_html=True,
    )

    fig_close = go.Figure()
    for col, name, color in [
        ("Cierre Final %", "Cierre Final %", GREEN),
        ("ICO a Corte %", "ICO a Corte %", SECONDARY_BLUE),
        ("Pendientes %", "Pendientes %", AMBER),
    ]:
        fig_close.add_trace(
            go.Scatter(
                x=op_period["PERIODO"],
                y=op_period[col] * 100,
                mode="lines+markers",
                name=name,
                line=dict(color=color, width=3),
                marker=dict(size=8),
                hovertemplate=f"%{{x}}<br>{name}: %{{y:.2f}}%<extra></extra>",
            )
        )

    fig_close.update_layout(
        height=360,
        margin=dict(l=20, r=20, t=70, b=35),
        legend=dict(
            orientation="h",
            y=1.16,
            x=0,
            font=dict(color="#F8FAFC", size=12),
        ),
        yaxis=dict(title="%", ticksuffix="%", rangemode="tozero"),
        hovermode="x unified",
    )
    st.plotly_chart(
        fig_close,
        use_container_width=True,
        key="story_closure_period",
    )


# =========================================================
# 3. SEGUNDA DIAGONAL · ¿DÓNDE ESTÁ EL PROBLEMA?
# =========================================================
st.markdown("## 3. ¿Dónde está el problema?")
st.caption(
    "Primero se identifica la composición del estado; después se localiza la concentración por cliente."
)

p1, p2 = st.columns([0.88, 1.12], gap="large")

with p1:
    st.markdown('<div class="question-tag">Composición general</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-title">ESTADO OPERATIVO</div>',
        unsafe_allow_html=True,
    )

    state_summary = pd.DataFrame(
        {
            "Estado": [
                "CUMPLIDO",
                "CUMPLIDO OPERATIVO",
                "EN PROGRAMACION",
                "EN TRANSITO",
            ],
            "Servicios": [
                cumplidos,
                cumpl_oper,
                en_programacion,
                en_transito,
            ],
        }
    )
    state_summary["Participación %"] = np.where(
        prod != 0,
        state_summary["Servicios"] / prod,
        np.nan,
    )

    fig_state = go.Figure(
        go.Bar(
            x=state_summary["Servicios"],
            y=state_summary["Estado"],
            orientation="h",
            marker_color=[
                GREEN,
                SECONDARY_BLUE,
                AMBER,
                ORANGE,
            ],
            text=state_summary["Servicios"].map(fmt_int),
            textposition="outside",
            hovertemplate="<b>%{y}</b><br>Servicios: %{x:,}<extra></extra>",
        )
    )
    fig_state.update_layout(
        height=360,
        margin=dict(l=10, r=60, t=25, b=25),
        yaxis=dict(autorange="reversed"),
        xaxis=dict(title="Servicios"),
        showlegend=False,
    )
    st.plotly_chart(
        fig_state,
        use_container_width=True,
        key="story_state_general",
    )

with p2:
    st.markdown('<div class="question-tag">Concentración por cliente</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-title">CLIENTES CON MAYOR PENDIENTE OPERATIVO</div>',
        unsafe_allow_html=True,
    )

    pending_clients = (
        client_perf.sort_values(
            ["Pendientes", "Servicios"],
            ascending=[False, False],
        )
        .head(10)
        .rename(columns={"CLIENTE": "Cliente"})
    )

    st.dataframe(
        pending_clients[
            [
                "Cliente", "Servicios", "Pendientes",
                "Pendientes %", "Cierre %"
            ]
        ].style.format(
            {
                "Servicios": "{:,.0f}",
                "Pendientes": "{:,.0f}",
                "Pendientes %": "{:.2%}",
                "Cierre %": "{:.2%}",
            }
        ),
        use_container_width=True,
        hide_index=True,
        height=360,
    )


# =========================================================
# 4. TERCERA DIAGONAL · ¿QUIÉN LO EXPLICA?
# =========================================================
st.markdown("## 4. ¿Quién explica el resultado?")
st.caption(
    "Se contrasta la concentración externa por cliente con la gestión interna de quien creó."
)

q1, q2 = st.columns([1.05, 0.95], gap="large")

with q1:
    st.markdown('<div class="question-tag">Cliente</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-title">DESEMPEÑO OPERATIVO POR CLIENTE</div>',
        unsafe_allow_html=True,
    )

    client_explain = (
        client_perf.sort_values(
            ["Servicios", "Pendientes"],
            ascending=[False, False],
        )
        .head(12)
        .rename(columns={"CLIENTE": "Cliente"})
    )

    st.dataframe(
        client_explain[
            [
                "Cliente", "Servicios", "Cumplidos",
                "Cumplido_Operativo", "Pendientes",
                "Cierre %", "Pendientes %"
            ]
        ].style.format(
            {
                "Servicios": "{:,.0f}",
                "Cumplidos": "{:,.0f}",
                "Cumplido_Operativo": "{:,.0f}",
                "Pendientes": "{:,.0f}",
                "Cierre %": "{:.2%}",
                "Pendientes %": "{:.2%}",
            }
        ),
        use_container_width=True,
        hide_index=True,
        height=430,
    )

with q2:
    st.markdown('<div class="question-tag">Gestión interna</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-title">DESEMPEÑO POR QUIÉN CREÓ</div>',
        unsafe_allow_html=True,
    )

    coord_explain = (
        coord_perf.sort_values(
            ["Servicios", "Pendientes"],
            ascending=[False, False],
        )
        .head(12)
        .rename(columns={"Quien Creó": "Quién creó"})
    )

    st.dataframe(
        coord_explain[
            [
                "Quién creó", "Servicios", "Cumplidos",
                "Cumplido_Operativo", "Pendientes",
                "Cierre %", "Pendientes %"
            ]
        ].style.format(
            {
                "Servicios": "{:,.0f}",
                "Cumplidos": "{:,.0f}",
                "Cumplido_Operativo": "{:,.0f}",
                "Pendientes": "{:,.0f}",
                "Cierre %": "{:.2%}",
                "Pendientes %": "{:.2%}",
            }
        ),
        use_container_width=True,
        hide_index=True,
        height=430,
    )


# =========================================================
# 5. ¿CON QUÉ RECURSOS OPERAMOS?
# =========================================================
st.markdown("## 5. ¿Con qué recursos estamos operando?")
st.caption(
    "La lectura de recursos se mantiene operacional: volumen y utilización, sin mezclar valores financieros."
)

res1, res2 = st.columns([0.8, 1.2], gap="large")

with res1:
    st.markdown(
        '<div class="section-title">FLOTA PROPIA VS TERCEROS · SERVICIOS</div>',
        unsafe_allow_html=True,
    )

    fleet_count = (
        valid["Tipo Flota"]
        .value_counts(dropna=False)
        .rename_axis("Tipo Flota")
        .reset_index(name="Servicios")
    )

    fig_fleet = go.Figure(
        go.Pie(
            labels=fleet_count["Tipo Flota"],
            values=fleet_count["Servicios"],
            hole=.62,
            textinfo="percent",
            marker=dict(
                colors=[ROYAL_BLUE, SECONDARY_BLUE, "#64748B"]
            ),
            hovertemplate="%{label}<br>%{value:,} servicios<br>%{percent}<extra></extra>",
        )
    )
    fig_fleet.update_layout(
        height=350,
        margin=dict(l=5, r=5, t=25, b=45),
        legend=dict(
            orientation="h",
            y=-0.07,
            x=0,
            font=dict(color="#F8FAFC", size=11),
        ),
    )
    st.plotly_chart(
        fig_fleet,
        use_container_width=True,
        key="story_fleet_mix",
    )

with res2:
    st.markdown(
        '<div class="section-title">SERVICIOS POR TIPOLOGÍA</div>',
        unsafe_allow_html=True,
    )

    typology = (
        valid.groupby("TIPOLOGIA", dropna=False)
        .size()
        .rename("Servicios")
        .reset_index()
        .sort_values("Servicios", ascending=False)
        .head(12)
    )

    fig_typology = go.Figure(
        go.Bar(
            x=typology["Servicios"],
            y=typology["TIPOLOGIA"],
            orientation="h",
            marker_color=SECONDARY_BLUE,
            text=typology["Servicios"].map(fmt_int),
            textposition="outside",
        )
    )
    fig_typology.update_layout(
        height=350,
        margin=dict(l=10, r=55, t=25, b=25),
        yaxis=dict(autorange="reversed"),
        xaxis=dict(title="Servicios"),
    )
    st.plotly_chart(
        fig_typology,
        use_container_width=True,
        key="story_typology",
    )

plate_ops = (
    valid.groupby("PLACA", dropna=False)
    .agg(
        Servicios=("CLIENTE", "size"),
        Clientes=("CLIENTE", "nunique"),
    )
    .reset_index()
    .sort_values(["Servicios", "Clientes"], ascending=[False, False])
    .head(15)
)
plate_ops.columns = ["Placa", "Servicios", "Clientes atendidos"]

driver_ops = (
    valid.groupby("CONDUCTOR_NOMBRE", dropna=False)
    .agg(
        Servicios=("CLIENTE", "size"),
        Clientes=("CLIENTE", "nunique"),
    )
    .reset_index()
    .sort_values(["Servicios", "Clientes"], ascending=[False, False])
    .head(15)
)
driver_ops.columns = ["Conductor", "Servicios", "Clientes atendidos"]

res3, res4 = st.columns(2, gap="large")

with res3:
    st.markdown(
        '<div class="section-title">TOP PLACAS POR SERVICIOS</div>',
        unsafe_allow_html=True,
    )
    st.dataframe(
        plate_ops.style.format(
            {
                "Servicios": "{:,.0f}",
                "Clientes atendidos": "{:,.0f}",
            }
        ),
        use_container_width=True,
        hide_index=True,
        height=390,
    )

with res4:
    st.markdown(
        '<div class="section-title">TOP CONDUCTORES POR SERVICIOS</div>',
        unsafe_allow_html=True,
    )
    st.dataframe(
        driver_ops.style.format(
            {
                "Servicios": "{:,.0f}",
                "Clientes atendidos": "{:,.0f}",
            }
        ),
        use_container_width=True,
        hide_index=True,
        height=390,
    )


# =========================================================
# 6. CIERRE DE LA HISTORIA · FOCOS DE ATENCIÓN
# =========================================================
st.markdown("## 6. ¿Dónde debemos actuar? · Focos de atención")
st.caption(
    "El dashboard termina priorizando los puntos que requieren revisión, en lugar de dejar al usuario buscando entre tablas."
)

# Cliente con más pendientes.
if not client_perf.empty:
    fc_client = (
        client_perf.sort_values(
            ["Pendientes", "Servicios"],
            ascending=[False, False],
        )
        .iloc[0]
    )
    focus_client_name = str(fc_client["CLIENTE"])
    focus_client_value = fmt_int(fc_client["Pendientes"])
    focus_client_note = f"{fmt_pct(fc_client['Pendientes %'])} de sus servicios"
else:
    focus_client_name = "Sin dato"
    focus_client_value = "0"
    focus_client_note = "Sin registros"

# Coordinador con menor cierre entre quienes tienen volumen material.
min_volume = max(10, int(prod * 0.003))
coord_material = coord_perf.loc[
    coord_perf["Servicios"] >= min_volume
].copy()
if coord_material.empty:
    coord_material = coord_perf.copy()

if not coord_material.empty:
    fc_coord = (
        coord_material.sort_values(
            ["Cierre %", "Servicios"],
            ascending=[True, False],
        )
        .iloc[0]
    )
    focus_coord_name = str(fc_coord["Quien Creó"])
    focus_coord_value = fmt_pct(fc_coord["Cierre %"])
    focus_coord_note = f"{fmt_int(fc_coord['Servicios'])} servicios"
else:
    focus_coord_name = "Sin dato"
    focus_coord_value = "—"
    focus_coord_note = "Sin registros"

# Cliente con mayor volumen en Cumplido Operativo.
if not client_perf.empty:
    fc_legal = (
        client_perf.sort_values(
            ["Cumplido_Operativo", "Servicios"],
            ascending=[False, False],
        )
        .iloc[0]
    )
    focus_legal_name = str(fc_legal["CLIENTE"])
    focus_legal_value = fmt_int(fc_legal["Cumplido_Operativo"])
    focus_legal_note = "servicios en legalización"
else:
    focus_legal_name = "Sin dato"
    focus_legal_value = "0"
    focus_legal_note = "Sin registros"

# Placa de mayor utilización.
if not plate_ops.empty:
    fc_plate = plate_ops.iloc[0]
    focus_plate_name = str(fc_plate["Placa"])
    focus_plate_value = fmt_int(fc_plate["Servicios"])
    focus_plate_note = f"{fmt_int(fc_plate['Clientes atendidos'])} clientes atendidos"
else:
    focus_plate_name = "Sin dato"
    focus_plate_value = "0"
    focus_plate_note = "Sin registros"

# Peor cierre por periodo: excluir el periodo en curso si está incompleto.
period_eval = op_period.copy()
if ultimo_periodo_incompleto and len(period_eval) > 1:
    period_eval = period_eval.iloc[:-1].copy()

if not period_eval.empty:
    fc_period = (
        period_eval.sort_values(
            ["Cierre Final %", "Servicios"],
            ascending=[True, False],
        )
        .iloc[0]
    )
    focus_period_name = str(fc_period["PERIODO"])
    focus_period_value = fmt_pct(fc_period["Cierre Final %"])
    focus_period_note = f"{fmt_int(fc_period['Servicios'])} servicios"
else:
    focus_period_name = "Sin dato"
    focus_period_value = "—"
    focus_period_note = "Sin periodo comparable"

st.markdown(
    f"""
    <div class="focus-grid">
        <div class="focus-card">
            <div class="label">Mayor concentración de pendientes</div>
            <div class="value">{focus_client_name}</div>
            <div class="note">{focus_client_value} pendientes · {focus_client_note}</div>
        </div>
        <div class="focus-card">
            <div class="label">Menor cierre con volumen material</div>
            <div class="value">{focus_coord_name}</div>
            <div class="note">{focus_coord_value} cierre · {focus_coord_note}</div>
        </div>
        <div class="focus-card">
            <div class="label">Mayor carga en legalización</div>
            <div class="value">{focus_legal_name}</div>
            <div class="note">{focus_legal_value} {focus_legal_note}</div>
        </div>
        <div class="focus-card">
            <div class="label">Placa con mayor utilización</div>
            <div class="value">{focus_plate_name}</div>
            <div class="note">{focus_plate_value} servicios · {focus_plate_note}</div>
        </div>
        <div class="focus-card">
            <div class="label">Periodo con menor cierre</div>
            <div class="value">{focus_period_name}</div>
            <div class="note">{focus_period_value} · {focus_period_note}</div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# 7. DETALLE PARA GESTIÓN
# =========================================================
st.markdown("## 7. Detalle para gestión")
st.caption(
    "Las tablas extensas quedan al final: sirven para profundizar después de entender la historia principal."
)


def build_status_table(
    data: pd.DataFrame,
    group_col: str,
    group_label: str,
) -> pd.DataFrame:
    x = valid_services(data).copy()

    states = [
        "CUMPLIDO",
        "CUMPLIDO OPERATIVO",
        "EN PROGRAMACION",
        "EN TRANSITO",
    ]

    pivot = (
        x.groupby([group_col, "ESTADO OP N"], dropna=False)
        .size()
        .unstack(fill_value=0)
    )

    for state in states:
        if state not in pivot.columns:
            pivot[state] = 0

    known = pivot[states].sum(axis=1)
    pivot["OTROS"] = pivot.sum(axis=1) - known
    pivot["Total general"] = pivot[states + ["OTROS"]].sum(axis=1)
    pivot["Pendientes operativos"] = (
        pivot["EN PROGRAMACION"] + pivot["EN TRANSITO"]
    )

    pivot["% Cumplido"] = np.where(
        pivot["Total general"].ne(0),
        pivot["CUMPLIDO"] / pivot["Total general"],
        np.nan,
    )
    pivot["% Cumplido Operativo"] = np.where(
        pivot["Total general"].ne(0),
        pivot["CUMPLIDO OPERATIVO"] / pivot["Total general"],
        np.nan,
    )
    pivot["Pendientes %"] = np.where(
        pivot["Total general"].ne(0),
        pivot["Pendientes operativos"] / pivot["Total general"],
        np.nan,
    )

    out = pivot[
        [
            "CUMPLIDO",
            "CUMPLIDO OPERATIVO",
            "EN PROGRAMACION",
            "EN TRANSITO",
            "OTROS",
            "Total general",
            "Pendientes operativos",
            "% Cumplido",
            "% Cumplido Operativo",
            "Pendientes %",
        ]
    ].copy()

    out = out.sort_values(
        ["Total general", "CUMPLIDO"],
        ascending=[False, False],
    )

    total_counts = out[
        [
            "CUMPLIDO",
            "CUMPLIDO OPERATIVO",
            "EN PROGRAMACION",
            "EN TRANSITO",
            "OTROS",
            "Total general",
            "Pendientes operativos",
        ]
    ].sum()

    total_general = total_counts["Total general"]

    # Convertir índice en columna visible: corrige QUIÉN CREÓ / CLIENTE.
    out = out.reset_index().rename(columns={group_col: group_label})

    total_row = pd.DataFrame(
        [
            {
                group_label: "TOTAL",
                "CUMPLIDO": total_counts["CUMPLIDO"],
                "CUMPLIDO OPERATIVO": total_counts["CUMPLIDO OPERATIVO"],
                "EN PROGRAMACION": total_counts["EN PROGRAMACION"],
                "EN TRANSITO": total_counts["EN TRANSITO"],
                "OTROS": total_counts["OTROS"],
                "Total general": total_general,
                "Pendientes operativos": total_counts["Pendientes operativos"],
                "% Cumplido": safe_div(
                    total_counts["CUMPLIDO"],
                    total_general,
                ),
                "% Cumplido Operativo": safe_div(
                    total_counts["CUMPLIDO OPERATIVO"],
                    total_general,
                ),
                "Pendientes %": safe_div(
                    total_counts["Pendientes operativos"],
                    total_general,
                ),
            }
        ]
    )

    return pd.concat([out, total_row], ignore_index=True)


with st.expander("📋 Estatus detallado por Quién creó / Cliente", expanded=False):
    status_mode = st.radio(
        "Ver tabla de estatus por",
        ["QUIÉN CREÓ", "CLIENTE"],
        horizontal=True,
        key="story_status_tracking_mode",
    )

    if status_mode == "QUIÉN CREÓ":
        status_tbl = build_status_table(
            valid,
            "Quien Creó",
            "Quién creó",
        )
    else:
        status_tbl = build_status_table(
            valid,
            "CLIENTE",
            "Cliente",
        )

    st.markdown(
        f'<div class="section-title">ESTATUS OPERATIVO POR {status_mode}</div>',
        unsafe_allow_html=True,
    )

    st.dataframe(
        status_tbl.style.format(
            {
                "CUMPLIDO": "{:,.0f}",
                "CUMPLIDO OPERATIVO": "{:,.0f}",
                "EN PROGRAMACION": "{:,.0f}",
                "EN TRANSITO": "{:,.0f}",
                "OTROS": "{:,.0f}",
                "Total general": "{:,.0f}",
                "Pendientes operativos": "{:,.0f}",
                "% Cumplido": "{:.2%}",
                "% Cumplido Operativo": "{:.2%}",
                "Pendientes %": "{:.2%}",
            }
        ),
        use_container_width=True,
        hide_index=True,
        height=540,
    )

# Garantizar status_tbl para descarga, aunque el expander no se abra.
status_tbl = build_status_table(
    valid,
    "Quien Creó",
    "Quién creó",
)

with st.expander("🧭 Matriz Cliente × Periodo", expanded=False):
    matrix_metric = st.radio(
        "Indicador de la matriz",
        [
            "Servicios",
            "Cierre Final %",
            "Cumplido Operativo %",
            "Pendientes %",
        ],
        horizontal=True,
        key="story_matrix_metric",
    )

    m = valid.copy()

    if grain == "Mensual":
        m["P_ORD"] = m["AÑO"] * 100 + m["MES_NUM"]
        m["P_LABEL"] = m.apply(
            lambda r: f"{MONTHS_ES.get(r['MES_NUM'], '')} {int(r['AÑO'])}",
            axis=1,
        )
    elif grain == "Bimestral":
        m["P_ORD"] = m["AÑO"] * 10 + m["BIMESTRE_NUM"].astype(int)
        m["P_LABEL"] = m.apply(
            lambda r: f"B{int(r['BIMESTRE_NUM'])} {int(r['AÑO'])}",
            axis=1,
        )
    elif grain == "Trimestral":
        m["P_ORD"] = m["AÑO"] * 10 + m["TRIMESTRE_NUM"].astype(int)
        m["P_LABEL"] = m.apply(
            lambda r: f"T{int(r['TRIMESTRE_NUM'])} {int(r['AÑO'])}",
            axis=1,
        )
    else:
        m["P_ORD"] = m["AÑO"] * 10 + m["SEMESTRE_NUM"].astype(int)
        m["P_LABEL"] = m.apply(
            lambda r: f"S{int(r['SEMESTRE_NUM'])} {int(r['AÑO'])}",
            axis=1,
        )

    period_order = (
        m[["P_ORD", "P_LABEL"]]
        .drop_duplicates()
        .sort_values("P_ORD")["P_LABEL"]
        .tolist()
    )

    m["_CUMPLIDO"] = m["ESTADO OP N"].eq("CUMPLIDO").astype(int)
    m["_CUMPLIDO_OPER"] = m["ESTADO OP N"].eq("CUMPLIDO OPERATIVO").astype(int)
    m["_PENDIENTE"] = m["ESTADO OP N"].isin(
        ["EN PROGRAMACION", "EN TRANSITO"]
    ).astype(int)

    matrix_base = (
        m.groupby(["CLIENTE", "P_LABEL"], dropna=False)
        .agg(
            Servicios=("CLIENTE", "size"),
            Cumplido=("_CUMPLIDO", "sum"),
            Cumplido_Operativo=("_CUMPLIDO_OPER", "sum"),
            Pendientes=("_PENDIENTE", "sum"),
        )
        .reset_index()
    )

    matrix_base["Cierre Final %"] = np.where(
        matrix_base["Servicios"].ne(0),
        matrix_base["Cumplido"] / matrix_base["Servicios"],
        np.nan,
    )
    matrix_base["Cumplido Operativo %"] = np.where(
        matrix_base["Servicios"].ne(0),
        matrix_base["Cumplido_Operativo"] / matrix_base["Servicios"],
        np.nan,
    )
    matrix_base["Pendientes %"] = np.where(
        matrix_base["Servicios"].ne(0),
        matrix_base["Pendientes"] / matrix_base["Servicios"],
        np.nan,
    )

    pivot = (
        matrix_base.pivot(
            index="CLIENTE",
            columns="P_LABEL",
            values=matrix_metric,
        )
        .reindex(columns=period_order)
    )

    client_order = (
        matrix_base.groupby("CLIENTE")["Servicios"]
        .sum()
        .sort_values(ascending=False)
        .index
    )
    pivot = pivot.reindex(client_order)

    # CLIENTE queda siempre visible como primera columna.
    pivot_display = (
        pivot.reset_index()
        .rename(columns={"CLIENTE": "Cliente"})
    )

    if matrix_metric == "Servicios":
        matrix_styled = pivot_display.style.format(
            {
                col: (lambda x: "" if pd.isna(x) else f"{x:,.0f}")
                for col in pivot_display.columns
                if col != "Cliente"
            }
        )
    else:
        matrix_styled = pivot_display.style.format(
            {
                col: "{:.2%}"
                for col in pivot_display.columns
                if col != "Cliente"
            }
        )

    st.dataframe(
        matrix_styled,
        use_container_width=True,
        hide_index=True,
        height=500,
    )

with st.expander("🚐 Servicios por placa y cliente", expanded=False):
    servicios_pc = (
        valid.assign(
            PLACA_N=valid["PLACA"].astype("string").str.strip().str.upper(),
            CLIENTE_N=valid["CLIENTE"].astype("string").str.strip().str.upper(),
        )
        .groupby(["PLACA_N", "CLIENTE_N"], dropna=False)
        .size()
        .rename("SERVICIOS")
        .reset_index()
        .rename(
            columns={
                "PLACA_N": "PLACA",
                "CLIENTE_N": "CLIENTE",
            }
        )
    )

    totales_placa = (
        servicios_pc.groupby("PLACA", as_index=False)["SERVICIOS"]
        .sum()
        .rename(columns={"SERVICIOS": "TOTAL_SERVICIOS_PLACA"})
    )

    servicios_pc = servicios_pc.merge(
        totales_placa,
        on="PLACA",
        how="left",
    )

    servicios_pc["PARTICIPACIÓN EN %"] = np.where(
        servicios_pc["TOTAL_SERVICIOS_PLACA"].ne(0),
        servicios_pc["SERVICIOS"] / servicios_pc["TOTAL_SERVICIOS_PLACA"],
        np.nan,
    )

    tabla_pc = (
        servicios_pc[
            ["PLACA", "CLIENTE", "SERVICIOS", "PARTICIPACIÓN EN %"]
        ]
        .sort_values(
            ["PLACA", "SERVICIOS", "CLIENTE"],
            ascending=[True, False, True],
        )
        .reset_index(drop=True)
    )

    st.dataframe(
        tabla_pc.style.format(
            {
                "SERVICIOS": "{:,.0f}",
                "PARTICIPACIÓN EN %": "{:.2%}",
            }
        ),
        use_container_width=True,
        hide_index=True,
        height=520,
    )

# Garantizar tabla_pc para la descarga aunque el expander no se abra.
servicios_pc = (
    valid.assign(
        PLACA_N=valid["PLACA"].astype("string").str.strip().str.upper(),
        CLIENTE_N=valid["CLIENTE"].astype("string").str.strip().str.upper(),
    )
    .groupby(["PLACA_N", "CLIENTE_N"], dropna=False)
    .size()
    .rename("SERVICIOS")
    .reset_index()
    .rename(columns={"PLACA_N": "PLACA", "CLIENTE_N": "CLIENTE"})
)
totales_placa = (
    servicios_pc.groupby("PLACA", as_index=False)["SERVICIOS"]
    .sum()
    .rename(columns={"SERVICIOS": "TOTAL_SERVICIOS_PLACA"})
)
servicios_pc = servicios_pc.merge(totales_placa, on="PLACA", how="left")
servicios_pc["PARTICIPACIÓN EN %"] = np.where(
    servicios_pc["TOTAL_SERVICIOS_PLACA"].ne(0),
    servicios_pc["SERVICIOS"] / servicios_pc["TOTAL_SERVICIOS_PLACA"],
    np.nan,
)
tabla_pc = (
    servicios_pc[
        ["PLACA", "CLIENTE", "SERVICIOS", "PARTICIPACIÓN EN %"]
    ]
    .sort_values(
        ["PLACA", "SERVICIOS", "CLIENTE"],
        ascending=[True, False, True],
    )
    .reset_index(drop=True)
)


# =========================================================
# 8. DESCARGA OPERATIVA
# =========================================================
st.markdown("## 8. Descarga operativa")

operational_excel = BytesIO()

with pd.ExcelWriter(operational_excel, engine="openpyxl") as writer:
    status_tbl.to_excel(
        writer,
        index=False,
        sheet_name="Estatus_Operativo",
    )
    client_perf.to_excel(
        writer,
        index=False,
        sheet_name="Clientes_Operativo",
    )
    coord_perf.to_excel(
        writer,
        index=False,
        sheet_name="Quien_Creo",
    )
    plate_ops.to_excel(
        writer,
        index=False,
        sheet_name="Placas",
    )
    driver_ops.to_excel(
        writer,
        index=False,
        sheet_name="Conductores",
    )
    tabla_pc.to_excel(
        writer,
        index=False,
        sheet_name="Placa_Cliente",
    )
    op_period.to_excel(
        writer,
        index=False,
        sheet_name="Evolucion_Operativa",
    )

    for ws in writer.book.worksheets:
        blue_fill = openpyxl.styles.PatternFill(
            fill_type="solid",
            fgColor="12365E",
        )
        white_bold = openpyxl.styles.Font(
            color="FFFFFF",
            bold=True,
        )
        for cell in ws[1]:
            cell.fill = blue_fill
            cell.font = white_bold
            cell.alignment = openpyxl.styles.Alignment(
                horizontal="center",
                vertical="center",
            )
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for col_cells in ws.columns:
            max_len = max(
                len(str(c.value)) if c.value is not None else 0
                for c in col_cells
            )
            ws.column_dimensions[col_cells[0].column_letter].width = min(
                max(max_len + 2, 12),
                44,
            )

operational_excel.seek(0)

st.download_button(
    "⬇️ Descargar Centro de Control Operativo",
    data=operational_excel,
    file_name="Centro_Control_Operativo_VSE.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True,
)
